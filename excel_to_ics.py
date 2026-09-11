"""Convert the editable timetable Excel workbook into an Apple-compatible ICS file."""

from __future__ import annotations

import argparse
import hashlib
import re
import sys
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta, timezone
from pathlib import Path
from typing import Any, Iterable
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel


PROJECT_DIR = Path(__file__).resolve().parent
DEFAULT_INPUT = PROJECT_DIR / "output" / "timetable.xlsx"
DEFAULT_OUTPUT = PROJECT_DIR / "output" / "timetable.ics"

EXPECTED_HEADERS = {
    "课程名称": "course_name",
    "星期": "weekday",
    "开始节次": "start_section",
    "结束节次": "end_section",
    "开始周": "start_week",
    "结束周": "end_week",
    "周期": "week_pattern",
    "指定周": "custom_weeks",
    "地点": "location",
    "教师": "teacher",
    "备注": "notes",
}
WEEK_PATTERNS = {"all", "odd", "even", "custom"}


class WorkbookFormatError(ValueError):
    """Fatal workbook-level error that prevents conversion."""


class CourseRowError(ValueError):
    """Error limited to one course row; other rows can still be converted."""


@dataclass(frozen=True)
class ConversionResult:
    output_path: Path
    event_count: int
    course_count: int
    skipped_course_count: int
    errors: tuple[str, ...]
    warnings: tuple[str, ...]


def _text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _integer(value: Any, label: str) -> int:
    if value is None or value == "":
        raise CourseRowError(f"{label} 不能为空")
    if isinstance(value, bool):
        raise CourseRowError(f"{label} 必须是整数")
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, str):
        try:
            return int(value.strip())
        except ValueError:
            pass
    raise CourseRowError(f"{label} 必须是整数")


def _setting_integer(value: Any, label: str) -> int:
    try:
        result = _integer(value, label)
    except CourseRowError as exc:
        raise WorkbookFormatError(str(exc)) from exc
    return result


def _date_value(value: Any, label: str) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        converted = from_excel(value)
        if isinstance(converted, datetime):
            return converted.date()
        if isinstance(converted, date):
            return converted
    if isinstance(value, str):
        text = value.strip()
        for format_string in ("%Y-%m-%d", "%Y/%m/%d"):
            try:
                return datetime.strptime(text, format_string).date()
            except ValueError:
                continue
    raise WorkbookFormatError(f"{label} 必须是日期，例如 2026-09-14")


def _time_value(value: Any, label: str) -> time:
    if isinstance(value, datetime):
        return value.time().replace(tzinfo=None)
    if isinstance(value, time):
        return value.replace(tzinfo=None)
    if isinstance(value, timedelta):
        total_seconds = int(value.total_seconds()) % (24 * 60 * 60)
        return time(total_seconds // 3600, (total_seconds % 3600) // 60, total_seconds % 60)
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        converted = from_excel(value)
        if isinstance(converted, datetime):
            return converted.time().replace(tzinfo=None)
        if isinstance(converted, time):
            return converted.replace(tzinfo=None)
    if isinstance(value, str):
        text = value.strip()
        for format_string in ("%H:%M", "%H:%M:%S"):
            try:
                return datetime.strptime(text, format_string).time()
            except ValueError:
                continue
    raise CourseRowError(f"{label} 必须是时间，例如 08:00")


def _read_settings(sheet: Any) -> dict[str, Any]:
    settings: dict[str, Any] = {}
    duplicates: list[str] = []
    for row_number, (key_cell, value_cell) in enumerate(
        sheet.iter_rows(min_row=2, max_col=2, values_only=True), start=2
    ):
        key = _text(key_cell)
        if not key:
            continue
        if key in settings:
            duplicates.append(f"Settings 第 {row_number} 行：参数 {key} 重复")
        settings[key] = value_cell
    if duplicates:
        raise WorkbookFormatError("；".join(duplicates))
    return settings


def _required_setting(settings: dict[str, Any], key: str) -> Any:
    if key not in settings or settings[key] in (None, ""):
        raise WorkbookFormatError(f"Settings 缺少 {key}")
    return settings[key]


def _read_course_headers(sheet: Any) -> dict[str, int]:
    header_positions: dict[str, int] = {}
    for column_index, cell in enumerate(sheet[1], start=1):
        label = _text(cell.value)
        if label in EXPECTED_HEADERS:
            header_positions[EXPECTED_HEADERS[label]] = column_index
    missing = [label for label, field in EXPECTED_HEADERS.items() if field not in header_positions]
    if missing:
        raise WorkbookFormatError(f"Courses 缺少列：{', '.join(missing)}")
    return header_positions


def _parse_custom_weeks(value: Any) -> list[int]:
    if value is None or value == "":
        return []
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        values = [value]
    else:
        values = re.split(r"[,，、\s]+", _text(value))
    weeks: list[int] = []
    for item in values:
        if item == "":
            continue
        week = _integer(item, "指定周")
        if week < 1:
            raise CourseRowError("指定周必须是大于 0 的整数")
        weeks.append(week)
    return sorted(set(weeks))


def _course_from_row(row: tuple[Any, ...], positions: dict[str, int]) -> dict[str, Any]:
    def value(field: str) -> Any:
        return row[positions[field] - 1] if positions[field] - 1 < len(row) else None

    course_name = _text(value("course_name"))
    if not course_name:
        raise CourseRowError("课程名称不能为空")

    course = {
        "course_name": course_name,
        "weekday": _integer(value("weekday"), "weekday"),
        "start_section": _integer(value("start_section"), "开始节次"),
        "end_section": _integer(value("end_section"), "结束节次"),
        "start_week": _integer(value("start_week"), "开始周"),
        "end_week": _integer(value("end_week"), "结束周"),
        "week_pattern": _text(value("week_pattern")).lower(),
        "custom_weeks": _parse_custom_weeks(value("custom_weeks")),
        "location": _text(value("location")),
        "teacher": _text(value("teacher")),
        "notes": _text(value("notes")),
    }

    if not 1 <= course["weekday"] <= 7:
        raise CourseRowError(f"weekday = {course['weekday']}，不合法（应为 1-7）")
    if course["start_section"] < 1 or course["end_section"] < 1:
        raise CourseRowError("节次必须大于 0")
    if course["start_section"] > course["end_section"]:
        raise CourseRowError("开始节次不能大于结束节次")
    if course["start_week"] < 1 or course["end_week"] < 1:
        raise CourseRowError("周次必须大于 0")
    if course["start_week"] > course["end_week"]:
        raise CourseRowError("开始周不能大于结束周")
    if course["week_pattern"] not in WEEK_PATTERNS:
        raise CourseRowError("周期必须是 all、odd、even 或 custom")
    if course["week_pattern"] == "custom":
        if not course["custom_weeks"]:
            raise CourseRowError("custom 周期必须填写指定周")
        outside = [
            week
            for week in course["custom_weeks"]
            if not course["start_week"] <= week <= course["end_week"]
        ]
        if outside:
            raise CourseRowError(
                f"指定周 {','.join(map(str, outside))} 不在开始周与结束周范围内"
            )
    return course


def _week_numbers(course: dict[str, Any]) -> list[int]:
    pattern = course["week_pattern"]
    if pattern == "custom":
        return course["custom_weeks"]
    weeks = list(range(course["start_week"], course["end_week"] + 1))
    if pattern == "odd":
        return [week for week in weeks if week % 2 == 1]
    if pattern == "even":
        return [week for week in weeks if week % 2 == 0]
    return weeks


def _section_time(settings: dict[str, Any], section: int, boundary: str) -> time:
    key = f"section_{section}_{boundary}"
    if key not in settings or settings[key] in (None, ""):
        raise CourseRowError(f"找不到 {key}")
    return _time_value(settings[key], key)


def _escape_ical(value: str) -> str:
    return (
        value.replace("\\", "\\\\")
        .replace("\r\n", "\n")
        .replace("\r", "\n")
        .replace("\n", "\\n")
        .replace(";", "\\;")
        .replace(",", "\\,")
    )


def _fold_ical_line(line: str) -> str:
    """Fold an iCalendar line at 75 UTF-8 octets without splitting characters."""
    chunks: list[str] = []
    current = ""
    for character in line:
        candidate = current + character
        if current and len(candidate.encode("utf-8")) > 75:
            chunks.append(current)
            current = " " + character
        else:
            current = candidate
    chunks.append(current)
    return "\r\n".join(chunks)


def _utc_ical(value: datetime) -> str:
    return value.astimezone(timezone.utc).strftime("%Y%m%dT%H%M%SZ")


def _event_uid(
    row_number: int, course: dict[str, Any], event_date: date, start_time: time
) -> str:
    stable_key = "|".join(
        (
            str(row_number),
            course["course_name"],
            event_date.isoformat(),
            start_time.isoformat(),
            course["location"],
        )
    )
    digest = hashlib.sha256(stable_key.encode("utf-8")).hexdigest()[:32]
    return f"{digest}@timetable-calendar.local"


def _event_lines(
    *,
    row_number: int,
    course: dict[str, Any],
    event_date: date,
    start_time: time,
    end_time: time,
    local_zone: ZoneInfo,
    alarm_minutes: int,
    timestamp: datetime,
) -> list[str]:
    local_start = datetime.combine(event_date, start_time, tzinfo=local_zone)
    local_end = datetime.combine(event_date, end_time, tzinfo=local_zone)
    description_parts = []
    if course["teacher"]:
        description_parts.append(f"教师：{course['teacher']}")
    if course["notes"]:
        description_parts.append(f"备注：{course['notes']}")

    lines = [
        "BEGIN:VEVENT",
        f"UID:{_event_uid(row_number, course, event_date, start_time)}",
        f"DTSTAMP:{_utc_ical(timestamp)}",
        f"DTSTART:{_utc_ical(local_start)}",
        f"DTEND:{_utc_ical(local_end)}",
        f"SUMMARY:{_escape_ical(course['course_name'])}",
    ]
    if course["location"]:
        lines.append(f"LOCATION:{_escape_ical(course['location'])}")
    if description_parts:
        lines.append(f"DESCRIPTION:{_escape_ical(chr(10).join(description_parts))}")
    lines.extend(("STATUS:CONFIRMED", "TRANSP:OPAQUE", "SEQUENCE:0"))
    if alarm_minutes > 0:
        lines.extend(
            (
                "BEGIN:VALARM",
                f"TRIGGER:-PT{alarm_minutes}M",
                "ACTION:DISPLAY",
                f"DESCRIPTION:{_escape_ical('课程提醒：' + course['course_name'])}",
                "END:VALARM",
            )
        )
    lines.append("END:VEVENT")
    return lines


def _serialize_calendar(lines: Iterable[str]) -> bytes:
    folded = [_fold_ical_line(line) for line in lines]
    return ("\r\n".join(folded) + "\r\n").encode("utf-8")


def convert_excel_to_ics(excel_path: Path, output_path: Path) -> ConversionResult:
    """Read Courses/Settings, skip bad course rows, and write one VEVENT per class date."""
    excel_path = excel_path.resolve()
    if not excel_path.is_file():
        raise WorkbookFormatError(f"找不到 Excel 文件：{excel_path}")

    workbook = load_workbook(excel_path, data_only=True)
    try:
        if "Courses" not in workbook.sheetnames:
            raise WorkbookFormatError("Excel 缺少 Courses Sheet")
        if "Settings" not in workbook.sheetnames:
            raise WorkbookFormatError("Excel 缺少 Settings Sheet")

        courses_sheet = workbook["Courses"]
        settings = _read_settings(workbook["Settings"])
        positions = _read_course_headers(courses_sheet)

        first_week_monday = _date_value(
            _required_setting(settings, "first_week_monday"), "first_week_monday"
        )
        if first_week_monday.weekday() != 0:
            raise WorkbookFormatError("first_week_monday 必须是周一")
        semester_end = _date_value(_required_setting(settings, "semester_end"), "semester_end")
        if semester_end < first_week_monday:
            raise WorkbookFormatError("semester_end 不能早于 first_week_monday")

        timezone_name = _text(settings.get("timezone", "Asia/Shanghai")) or "Asia/Shanghai"
        try:
            local_zone = ZoneInfo(timezone_name)
        except ZoneInfoNotFoundError as exc:
            raise WorkbookFormatError(f"无法识别时区：{timezone_name}") from exc

        alarm_minutes = _setting_integer(settings.get("default_alarm_minutes", 15), "default_alarm_minutes")
        if alarm_minutes < 0:
            raise WorkbookFormatError("default_alarm_minutes 不能小于 0")
        semester_name = _text(settings.get("semester_name", "课程表")) or "课程表"

        timestamp = datetime.now(timezone.utc).replace(microsecond=0)
        calendar_lines = [
            "BEGIN:VCALENDAR",
            "VERSION:2.0",
            "PRODID:-//Timetable Calendar V1//ZH-CN",
            "CALSCALE:GREGORIAN",
            "METHOD:PUBLISH",
            f"X-WR-CALNAME:{_escape_ical(semester_name)}",
            f"X-WR-TIMEZONE:{_escape_ical(timezone_name)}",
        ]
        errors: list[str] = []
        warnings: list[str] = []
        event_count = 0
        course_count = 0
        skipped_course_count = 0

        max_column = max(positions.values())
        for row_number, row in enumerate(
            courses_sheet.iter_rows(min_row=2, max_col=max_column, values_only=True), start=2
        ):
            if all(value in (None, "") for value in row):
                continue
            course_count += 1
            try:
                course = _course_from_row(row, positions)
                start_time = _section_time(settings, course["start_section"], "start")
                end_time = _section_time(settings, course["end_section"], "end")
                if end_time <= start_time:
                    raise CourseRowError(
                        f"section_{course['end_section']}_end 必须晚于 "
                        f"section_{course['start_section']}_start"
                    )

                row_event_count = 0
                for week_number in _week_numbers(course):
                    event_date = first_week_monday + timedelta(
                        days=(week_number - 1) * 7 + course["weekday"] - 1
                    )
                    if event_date > semester_end:
                        warnings.append(
                            f"第 {row_number} 行：第 {week_number} 周日期 {event_date} "
                            "晚于 semester_end，已跳过"
                        )
                        continue
                    calendar_lines.extend(
                        _event_lines(
                            row_number=row_number,
                            course=course,
                            event_date=event_date,
                            start_time=start_time,
                            end_time=end_time,
                            local_zone=local_zone,
                            alarm_minutes=alarm_minutes,
                            timestamp=timestamp,
                        )
                    )
                    event_count += 1
                    row_event_count += 1
                if row_event_count == 0:
                    skipped_course_count += 1
                    errors.append(f"第 {row_number} 行：没有可生成的上课日期")
            except CourseRowError as exc:
                skipped_course_count += 1
                errors.append(f"第 {row_number} 行：{exc}")

        calendar_lines.append("END:VCALENDAR")
        if event_count == 0:
            detail = "；".join(errors) if errors else "Courses 中没有课程"
            raise WorkbookFormatError(f"没有生成任何日历事件：{detail}")

        output_path = output_path.resolve()
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_bytes(_serialize_calendar(calendar_lines))
        return ConversionResult(
            output_path=output_path,
            event_count=event_count,
            course_count=course_count,
            skipped_course_count=skipped_course_count,
            errors=tuple(errors),
            warnings=tuple(warnings),
        )
    finally:
        workbook.close()


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Excel 课表转 Apple Calendar ICS")
    parser.add_argument("excel", nargs="?", type=Path, default=DEFAULT_INPUT, help="课表 Excel")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT, help="ICS 输出路径")
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    try:
        result = convert_excel_to_ics(args.excel, args.output)
    except Exception as exc:  # CLI boundary: keep errors concise for new Python users.
        print(f"错误：{exc}", file=sys.stderr)
        return 1

    print(f"已读取 {result.course_count} 行课程，生成 {result.event_count} 个日历事件")
    print(f"ICS：{result.output_path}")
    for warning in result.warnings:
        print(f"警告：{warning}", file=sys.stderr)
    for error in result.errors:
        print(f"跳过：{error}", file=sys.stderr)
    if result.errors:
        print(f"共有 {result.skipped_course_count} 行课程未生成事件，请修正后重试。", file=sys.stderr)
        return 2
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

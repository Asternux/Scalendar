"""Convert a timetable image into a normalized JSON document and an Excel workbook.

The image is sent to an OpenAI vision-capable model.  The workbook is the
editable source of truth for the second step, ``excel_to_ics.py``.
"""

from __future__ import annotations

import argparse
import base64
import json
import mimetypes
import os
import re
import sys
from datetime import date, time
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.worksheet.datavalidation import DataValidation


PROJECT_DIR = Path(__file__).resolve().parent
DEFAULT_IMAGE = PROJECT_DIR / "input" / "timetable.png"
DEFAULT_OUTPUT = PROJECT_DIR / "output" / "timetable.xlsx"
DEFAULT_RAW_RESPONSE = PROJECT_DIR / "temp" / "raw_response.txt"
DEFAULT_NORMALIZED_JSON = PROJECT_DIR / "temp" / "normalized_timetable.json"
DEFAULT_MODEL = "gpt-5.6-terra"

COURSE_FIELDS = (
    "course_name",
    "weekday",
    "start_section",
    "end_section",
    "start_week",
    "end_week",
    "week_pattern",
    "custom_weeks",
    "location",
    "teacher",
    "notes",
)

WEEK_PATTERNS = {"all", "odd", "even", "custom"}

TIMETABLE_JSON_SCHEMA: dict[str, Any] = {
    "type": "object",
    "additionalProperties": False,
    "required": ["courses"],
    "properties": {
        "courses": {
            "type": "array",
            "items": {
                "type": "object",
                "additionalProperties": False,
                "required": list(COURSE_FIELDS),
                "properties": {
                    "course_name": {"type": "string"},
                    "weekday": {"type": ["integer", "null"]},
                    "start_section": {"type": ["integer", "null"]},
                    "end_section": {"type": ["integer", "null"]},
                    "start_week": {"type": ["integer", "null"]},
                    "end_week": {"type": ["integer", "null"]},
                    "week_pattern": {
                        "type": "string",
                        "enum": ["all", "odd", "even", "custom"],
                    },
                    "custom_weeks": {
                        "type": "array",
                        "items": {"type": "integer"},
                    },
                    "location": {"type": "string"},
                    "teacher": {"type": "string"},
                    "notes": {"type": "string"},
                },
            },
        }
    },
}

VISION_INSTRUCTIONS = """你负责把大学课表图片转换为结构化数据。请理解整张表的行列、星期和合并单元格关系，不要只做逐字 OCR。

规则：
1. weekday 使用 1=周一 到 7=周日。
2. 识别课程名、开始/结束节次、开始/结束周、单双周、指定周、地点、教师和备注。
3. 普通周次使用 week_pattern=all；单周使用 odd；双周使用 even；明确列举周次时使用 custom。
4. custom_weeks 始终返回数组；非 custom 时返回空数组。
5. 不确定的字符串填空字符串，不确定的数字填 null。不要根据常识猜测图片里看不清的内容。
6. 同一课程在不同星期、节次、地点或周次出现时，分别输出记录。
7. 仔细检查跨节课程，例如 3-4 节必须是 start_section=3、end_section=4。
8. 只返回符合给定 schema 的结果。"""


class TimetableValidationError(ValueError):
    """Raised when timetable JSON is structurally inconsistent."""

    def __init__(self, errors: list[str]):
        self.errors = errors
        super().__init__("\n".join(errors))


def _string(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _optional_int(value: Any, field_name: str, course_number: int) -> int | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        raise ValueError(f"第 {course_number} 门课程：{field_name} 必须是整数或 null")
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, str):
        try:
            return int(value.strip())
        except ValueError:
            pass
    raise ValueError(f"第 {course_number} 门课程：{field_name} 必须是整数或 null")


def _custom_weeks(value: Any, course_number: int) -> list[int]:
    if value is None or value == "":
        return []
    if isinstance(value, str):
        values: list[Any] = [part.strip() for part in value.replace("，", ",").split(",")]
    elif isinstance(value, list):
        values = value
    else:
        raise ValueError(f"第 {course_number} 门课程：custom_weeks 必须是整数数组")

    weeks: list[int] = []
    for item in values:
        week = _optional_int(item, "custom_weeks", course_number)
        if week is None or week < 1:
            raise ValueError(f"第 {course_number} 门课程：指定周必须是大于 0 的整数")
        weeks.append(week)
    return sorted(set(weeks))


def validate_and_normalize(data: Any) -> dict[str, list[dict[str, Any]]]:
    """Validate model JSON and return a predictable field order and value types."""
    if not isinstance(data, dict):
        raise TimetableValidationError(["JSON 顶层必须是 object"])
    if not isinstance(data.get("courses"), list):
        raise TimetableValidationError(["courses 必须是 list"])

    normalized_courses: list[dict[str, Any]] = []
    errors: list[str] = []

    for index, raw_course in enumerate(data["courses"], start=1):
        if not isinstance(raw_course, dict):
            errors.append(f"第 {index} 门课程必须是 object")
            continue

        try:
            course = {
                "course_name": _string(raw_course.get("course_name", "")),
                "weekday": _optional_int(raw_course.get("weekday"), "weekday", index),
                "start_section": _optional_int(
                    raw_course.get("start_section"), "start_section", index
                ),
                "end_section": _optional_int(
                    raw_course.get("end_section"), "end_section", index
                ),
                "start_week": _optional_int(raw_course.get("start_week"), "start_week", index),
                "end_week": _optional_int(raw_course.get("end_week"), "end_week", index),
                "week_pattern": _string(raw_course.get("week_pattern", "")),
                "custom_weeks": _custom_weeks(raw_course.get("custom_weeks", []), index),
                "location": _string(raw_course.get("location", "")),
                "teacher": _string(raw_course.get("teacher", "")),
                "notes": _string(raw_course.get("notes", "")),
            }

            weekday = course["weekday"]
            if weekday is not None and not 1 <= weekday <= 7:
                raise ValueError(f"第 {index} 门课程：weekday 必须在 1-7 之间")

            for field_name in ("start_section", "end_section", "start_week", "end_week"):
                value = course[field_name]
                if value is not None and value < 1:
                    raise ValueError(f"第 {index} 门课程：{field_name} 必须大于 0")

            if (
                course["start_section"] is not None
                and course["end_section"] is not None
                and course["start_section"] > course["end_section"]
            ):
                raise ValueError(f"第 {index} 门课程：start_section 不能大于 end_section")

            if (
                course["start_week"] is not None
                and course["end_week"] is not None
                and course["start_week"] > course["end_week"]
            ):
                raise ValueError(f"第 {index} 门课程：start_week 不能大于 end_week")

            pattern = course["week_pattern"]
            if pattern not in WEEK_PATTERNS:
                raise ValueError(
                    f"第 {index} 门课程：week_pattern 必须是 all、odd、even 或 custom"
                )
            if pattern == "custom" and not course["custom_weeks"]:
                raise ValueError(f"第 {index} 门课程：custom 模式必须提供 custom_weeks")
            if pattern != "custom":
                course["custom_weeks"] = []

            normalized_courses.append(course)
        except ValueError as exc:
            errors.append(str(exc))

    if errors:
        raise TimetableValidationError(errors)
    return {"courses": normalized_courses}


def _image_data_url(image_path: Path) -> str:
    mime_type, _ = mimetypes.guess_type(image_path.name)
    if not mime_type or not mime_type.startswith("image/"):
        mime_type = "image/png"
    encoded = base64.b64encode(image_path.read_bytes()).decode("ascii")
    return f"data:{mime_type};base64,{encoded}"


def request_timetable_json(
    image_path: Path,
    raw_response_path: Path,
    *,
    api_key: str | None = None,
    model: str | None = None,
) -> dict[str, Any]:
    """Call the Responses API with image input and strict Structured Outputs."""
    resolved_api_key = (api_key or os.getenv("OPENAI_API_KEY", "")).strip()
    if not resolved_api_key:
        raise RuntimeError("未检测到 OPENAI_API_KEY")

    try:
        from openai import OpenAI
    except ImportError as exc:
        raise RuntimeError("未安装 openai。请先运行：pip install -r requirements.txt") from exc

    resolved_model = (model or os.getenv("OPENAI_MODEL", DEFAULT_MODEL)).strip() or DEFAULT_MODEL
    client = OpenAI(api_key=resolved_api_key)
    response = client.responses.create(
        model=resolved_model,
        input=[
            {"role": "developer", "content": VISION_INSTRUCTIONS},
            {
                "role": "user",
                "content": [
                    {
                        "type": "input_text",
                        "text": "请读取这张课表图片，并输出所有课程记录。",
                    },
                    {
                        "type": "input_image",
                        "image_url": _image_data_url(image_path),
                        "detail": "high",
                    },
                ],
            },
        ],
        text={
            "format": {
                "type": "json_schema",
                "name": "university_timetable",
                "strict": True,
                "schema": TIMETABLE_JSON_SCHEMA,
            }
        },
    )

    raw_text = response.output_text or ""
    raw_response_path.parent.mkdir(parents=True, exist_ok=True)
    raw_response_path.write_text(raw_text, encoding="utf-8")
    if not raw_text.strip():
        raise RuntimeError("模型返回了空响应；原始响应已保存到 temp/raw_response.txt")

    try:
        return json.loads(raw_text)
    except json.JSONDecodeError as exc:
        raise RuntimeError(
            "模型响应不是有效 JSON；原始响应已保存到 temp/raw_response.txt："
            f"{exc}"
        ) from exc


def default_settings() -> list[tuple[str, Any]]:
    """Return editable semester and section settings with typed dates/times."""
    rows: list[tuple[str, Any]] = [
        ("semester_name", "2026秋季学期"),
        ("first_week_monday", date(2026, 9, 14)),
        ("semester_end", date(2027, 1, 15)),
        ("default_alarm_minutes", 15),
        ("timezone", "Asia/Shanghai"),
    ]
    section_times = {
        1: (time(8, 0), time(8, 45)),
        2: (time(8, 55), time(9, 40)),
        3: (time(10, 0), time(10, 45)),
        4: (time(10, 55), time(11, 40)),
        5: (time(14, 0), time(14, 45)),
        6: (time(14, 55), time(15, 40)),
        7: (time(16, 0), time(16, 45)),
        8: (time(16, 55), time(17, 40)),
        9: (time(19, 0), time(19, 45)),
        10: (time(19, 55), time(20, 40)),
        11: (time(20, 50), time(21, 35)),
        12: (time(21, 45), time(22, 30)),
    }
    for section, (start, end) in section_times.items():
        rows.append((f"section_{section}_start", start))
        rows.append((f"section_{section}_end", end))
    return rows


def load_existing_settings(workbook_path: Path) -> dict[str, Any]:
    """Read editable Settings values from an existing workbook, if present."""
    workbook_path = workbook_path.resolve()
    if not workbook_path.is_file():
        return {}
    workbook = load_workbook(workbook_path, data_only=True, read_only=True)
    try:
        if "Settings" not in workbook.sheetnames:
            return {}
        settings: dict[str, Any] = {}
        for key, value in workbook["Settings"].iter_rows(
            min_row=2, max_col=2, values_only=True
        ):
            normalized_key = _string(key)
            if normalized_key and value not in (None, ""):
                settings[normalized_key] = value
        return settings
    finally:
        workbook.close()


def _style_header(worksheet: Any, cell_range: str) -> None:
    header = worksheet[cell_range][0]
    fill = PatternFill("solid", fgColor="1F4E78")
    border = Border(bottom=Side(style="thin", color="D9E2F3"))
    for cell in header:
        cell.fill = fill
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border


def write_excel(
    timetable: dict[str, list[dict[str, Any]]],
    output_path: Path,
    settings_overrides: dict[str, Any] | None = None,
    *,
    preserve_existing_settings: bool = True,
    replace_settings: bool = False,
) -> Path:
    """Write normalized timetable data to a simple two-sheet Excel workbook."""
    output_path = output_path.resolve()
    setting_rows = default_settings()
    base_keys = [key for key, _ in setting_rows]
    settings_by_key = {} if replace_settings else dict(setting_rows)
    if preserve_existing_settings and not replace_settings:
        for key, value in load_existing_settings(output_path).items():
            settings_by_key[key] = value
    if settings_overrides:
        for key, value in settings_overrides.items():
            settings_by_key[key] = value

    def settings_sort_key(key: str) -> tuple[int, int | str, int]:
        if key in base_keys:
            return (0, base_keys.index(key), 0)
        match = re.fullmatch(r"section_(\d+)_(start|end)", key)
        if match:
            return (1, int(match.group(1)), 0 if match.group(2) == "start" else 1)
        return (2, key, 0)

    ordered_setting_keys = sorted(settings_by_key, key=settings_sort_key)

    workbook = Workbook()
    courses_sheet = workbook.active
    courses_sheet.title = "Courses"
    settings_sheet = workbook.create_sheet("Settings")

    course_headers = [
        "课程名称",
        "星期",
        "开始节次",
        "结束节次",
        "开始周",
        "结束周",
        "周期",
        "指定周",
        "地点",
        "教师",
        "备注",
    ]
    courses_sheet.append(course_headers)
    for course in timetable["courses"]:
        courses_sheet.append(
            [
                course["course_name"],
                course["weekday"],
                course["start_section"],
                course["end_section"],
                course["start_week"],
                course["end_week"],
                course["week_pattern"],
                ",".join(str(week) for week in course["custom_weeks"]),
                course["location"],
                course["teacher"],
                course["notes"],
            ]
        )

    _style_header(courses_sheet, "A1:K1")
    courses_sheet.freeze_panes = "A2"
    courses_sheet.auto_filter.ref = f"A1:K{max(1, courses_sheet.max_row)}"
    courses_sheet.sheet_view.showGridLines = False
    courses_sheet.row_dimensions[1].height = 24
    widths = [22, 9, 11, 11, 9, 9, 11, 18, 20, 16, 28]
    for column, width in zip("ABCDEFGHIJK", widths, strict=True):
        courses_sheet.column_dimensions[column].width = width
    for row in courses_sheet.iter_rows(min_row=2, max_row=courses_sheet.max_row):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="center", wrap_text=False)
            cell.protection = Protection(locked=False)
        for cell in row[1:8]:
            cell.alignment = Alignment(horizontal="center", vertical="center")

    weekday_validation = DataValidation(type="whole", operator="between", formula1="1", formula2="7")
    section_validation = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="1")
    week_validation = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="1")
    pattern_validation = DataValidation(type="list", formula1='"all,odd,even,custom"')
    courses_sheet.add_data_validation(weekday_validation)
    courses_sheet.add_data_validation(section_validation)
    courses_sheet.add_data_validation(week_validation)
    courses_sheet.add_data_validation(pattern_validation)
    weekday_validation.add("B2:B1000")
    section_validation.add("C2:D1000")
    week_validation.add("E2:F1000")
    pattern_validation.add("G2:G1000")
    invalid_fill = PatternFill("solid", fgColor="FCE8E6")
    invalid_font = Font(name="Arial", size=10, bold=True, color="B91C1C")
    courses_sheet.conditional_formatting.add(
        "B2:B1000", FormulaRule(formula=["=AND(B2<>\"\",OR(B2<1,B2>7))"], fill=invalid_fill, font=invalid_font)
    )

    settings_sheet.append(["参数", "值"])
    for key in ordered_setting_keys:
        settings_sheet.append([key, settings_by_key[key]])
    _style_header(settings_sheet, "A1:B1")
    settings_sheet.freeze_panes = "A2"
    settings_sheet.sheet_view.showGridLines = False
    settings_sheet.row_dimensions[1].height = 24
    settings_sheet.column_dimensions["A"].width = 28
    settings_sheet.column_dimensions["B"].width = 22
    input_fill = PatternFill("solid", fgColor="FFF2CC")
    for row in settings_sheet.iter_rows(min_row=2, max_row=settings_sheet.max_row):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="center")
        row[1].fill = input_fill
        row[1].protection = Protection(locked=False)
        key = row[0].value
        if key in {"first_week_monday", "semester_end"}:
            row[1].number_format = "yyyy-mm-dd"
        elif str(key).startswith("section_"):
            row[1].number_format = "hh:mm"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path


def _load_json_file(path: Path) -> dict[str, Any]:
    try:
        return json.loads(path.read_text(encoding="utf-8-sig"))
    except json.JSONDecodeError as exc:
        raise RuntimeError(f"JSON 文件解析失败：{exc}") from exc


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="课表图片转标准 JSON 与 Excel")
    parser.add_argument("image", nargs="?", type=Path, help="课表图片；默认 input/timetable.png")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT, help="Excel 输出路径")
    parser.add_argument(
        "--json",
        dest="json_path",
        type=Path,
        help="跳过 API，直接把已有标准 JSON 转成 Excel（便于测试或修正）",
    )
    parser.add_argument(
        "--raw-response",
        type=Path,
        default=DEFAULT_RAW_RESPONSE,
        help="模型原始响应保存路径",
    )
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    try:
        if args.json_path:
            json_path = args.json_path.resolve()
            if not json_path.is_file():
                raise RuntimeError(f"找不到 JSON 文件：{json_path}")
            raw_data = _load_json_file(json_path)
        else:
            image_path = (args.image or DEFAULT_IMAGE).resolve()
            if not image_path.is_file():
                raise RuntimeError(f"找不到课表图片：{image_path}")
            raw_data = request_timetable_json(image_path, args.raw_response.resolve())

        timetable = validate_and_normalize(raw_data)
        DEFAULT_NORMALIZED_JSON.parent.mkdir(parents=True, exist_ok=True)
        DEFAULT_NORMALIZED_JSON.write_text(
            json.dumps(timetable, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        output_path = write_excel(timetable, args.output)
        print(f"已识别 {len(timetable['courses'])} 门课程")
        print(f"规范化 JSON：{DEFAULT_NORMALIZED_JSON}")
        print(f"Excel：{output_path}")
        print("请先检查 Courses 和 Settings，再运行 excel_to_ics.py。")
        return 0
    except TimetableValidationError as exc:
        print("课表 JSON 验证失败：", file=sys.stderr)
        for error in exc.errors:
            print(f"- {error}", file=sys.stderr)
        return 1
    except Exception as exc:  # CLI boundary: present a concise, actionable message.
        print(f"错误：{exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())

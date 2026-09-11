from __future__ import annotations

import tempfile
import unittest
from datetime import date, time
from pathlib import Path

from openpyxl import load_workbook

from excel_to_ics import convert_excel_to_ics
from image_to_excel import TimetableValidationError, validate_and_normalize, write_excel
from app import apply_course_patch, calculate_first_week_monday, read_excel_for_gui


def sample_data() -> dict:
    return {
        "courses": [
            {
                "course_name": "线性代数",
                "weekday": 1,
                "start_section": 3,
                "end_section": 4,
                "start_week": 1,
                "end_week": 3,
                "week_pattern": "all",
                "custom_weeks": [],
                "location": "二教301",
                "teacher": "张老师",
                "notes": "带教材",
            },
            {
                "course_name": "体育",
                "weekday": 2,
                "start_section": 5,
                "end_section": 6,
                "start_week": 1,
                "end_week": 5,
                "week_pattern": "odd",
                "custom_weeks": [],
                "location": "操场",
                "teacher": "王老师",
                "notes": "",
            },
            {
                "course_name": "物理实验",
                "weekday": 3,
                "start_section": 1,
                "end_section": 2,
                "start_week": 2,
                "end_week": 6,
                "week_pattern": "even",
                "custom_weeks": [],
                "location": "实验楼A-201",
                "teacher": "李老师",
                "notes": "分组实验",
            },
            {
                "course_name": "学术写作",
                "weekday": 4,
                "start_section": 7,
                "end_section": 8,
                "start_week": 1,
                "end_week": 8,
                "week_pattern": "custom",
                "custom_weeks": [1, 3, 5, 8],
                "location": "图书馆研讨室",
                "teacher": "陈老师",
                "notes": "提交草稿",
            },
        ]
    }


class JsonValidationTests(unittest.TestCase):
    def test_normalizes_custom_weeks_and_missing_strings(self) -> None:
        data = {
            "courses": [
                {
                    "course_name": "测试课",
                    "weekday": "3",
                    "start_section": 1,
                    "end_section": 2,
                    "start_week": 1,
                    "end_week": 9,
                    "week_pattern": "custom",
                    "custom_weeks": "9, 1, 3, 3",
                }
            ]
        }
        course = validate_and_normalize(data)["courses"][0]
        self.assertEqual(course["weekday"], 3)
        self.assertEqual(course["custom_weeks"], [1, 3, 9])
        self.assertEqual(course["teacher"], "")

    def test_rejects_invalid_ranges_and_weekday(self) -> None:
        data = sample_data()
        data["courses"][0]["weekday"] = 9
        data["courses"][0]["start_section"] = 4
        data["courses"][0]["end_section"] = 3
        with self.assertRaises(TimetableValidationError) as context:
            validate_and_normalize(data)
        self.assertIn("weekday 必须在 1-7", str(context.exception))


class GuiLogicTests(unittest.TestCase):
    def test_reference_week_calculates_first_week(self) -> None:
        self.assertEqual(
            calculate_first_week_monday(date(2026, 9, 14), 3),
            date(2026, 8, 31),
        )

    def test_reference_date_must_be_monday(self) -> None:
        with self.assertRaisesRegex(ValueError, "必须是周一"):
            calculate_first_week_monday(date(2026, 9, 15), 3)

    def test_regeneration_preserves_existing_dates_and_section_times(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            xlsx_path = Path(temporary_directory) / "timetable.xlsx"
            normalized = validate_and_normalize(sample_data())
            write_excel(normalized, xlsx_path)
            workbook = load_workbook(xlsx_path)
            try:
                workbook["Settings"]["B3"] = date(2026, 8, 31)
                workbook["Settings"]["B7"] = time(7, 30)
                workbook.save(xlsx_path)
            finally:
                workbook.close()

            write_excel(normalized, xlsx_path)
            _courses, settings = read_excel_for_gui(xlsx_path)
            first_week = settings["first_week_monday"]
            if hasattr(first_week, "date"):
                first_week = first_week.date()
            self.assertEqual(first_week, date(2026, 8, 31))
            self.assertEqual(settings["section_1_start"], time(7, 30))

    def test_batch_course_edit_updates_only_selected_rows(self) -> None:
        courses = validate_and_normalize(sample_data())["courses"]
        updated = apply_course_patch(courses, [0, 2], {"location": "临时教室"})
        self.assertEqual(updated[0]["location"], "临时教室")
        self.assertEqual(updated[1]["location"], "操场")
        self.assertEqual(updated[2]["location"], "临时教室")

    def test_dynamic_section_count_is_written_without_default_rows(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            xlsx_path = Path(temporary_directory) / "dynamic_sections.xlsx"
            settings = {
                "semester_name": "测试学期",
                "first_week_monday": date(2026, 8, 31),
                "semester_end": date(2027, 1, 15),
                "default_alarm_minutes": 15,
                "timezone": "Asia/Shanghai",
            }
            for section in range(1, 15):
                start_minutes = 8 * 60 + (section - 1) * 50
                end_minutes = start_minutes + 45
                settings[f"section_{section}_start"] = time(
                    start_minutes // 60, start_minutes % 60
                )
                settings[f"section_{section}_end"] = time(
                    end_minutes // 60, end_minutes % 60
                )
            write_excel(
                validate_and_normalize(sample_data()),
                xlsx_path,
                settings_overrides=settings,
                preserve_existing_settings=False,
                replace_settings=True,
            )
            _courses, loaded_settings = read_excel_for_gui(xlsx_path)
            self.assertEqual(loaded_settings["section_14_start"], time(18, 50))
            self.assertNotIn("section_15_start", loaded_settings)


class EndToEndTests(unittest.TestCase):
    def test_json_to_excel_to_ics(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            xlsx_path = root / "timetable.xlsx"
            ics_path = root / "timetable.ics"
            normalized = validate_and_normalize(sample_data())
            write_excel(normalized, xlsx_path)

            workbook = load_workbook(xlsx_path, data_only=False)
            try:
                self.assertEqual(workbook.sheetnames, ["Courses", "Settings"])
                self.assertEqual(workbook["Courses"].freeze_panes, "A2")
                self.assertEqual(workbook["Courses"]["A2"].value, "线性代数")
                self.assertEqual(workbook["Courses"]["H5"].value, "1,3,5,8")
                self.assertEqual(workbook["Settings"]["A3"].value, "first_week_monday")
                self.assertEqual(workbook["Settings"]["B3"].number_format, "yyyy-mm-dd")
            finally:
                workbook.close()

            result = convert_excel_to_ics(xlsx_path, ics_path)
            self.assertEqual(result.course_count, 4)
            self.assertEqual(result.event_count, 13)
            self.assertEqual(result.skipped_course_count, 0)
            raw = ics_path.read_bytes()
            self.assertTrue(raw.endswith(b"\r\n"))
            text = raw.decode("utf-8")
            self.assertEqual(text.count("BEGIN:VEVENT"), 13)
            self.assertEqual(text.count("BEGIN:VALARM"), 13)
            self.assertNotIn("RRULE", text)
            self.assertIn("DTSTART:20260914T020000Z", text)
            self.assertIn("DTEND:20260914T034000Z", text)
            self.assertIn("DTSTART:20260915T060000Z", text)  # odd week 1
            self.assertIn("DTSTART:20260929T060000Z", text)  # odd week 3
            self.assertIn("DTSTART:20261013T060000Z", text)  # odd week 5
            self.assertIn("DTSTART:20260923T000000Z", text)  # even week 2
            self.assertIn("DTSTART:20261007T000000Z", text)  # even week 4
            self.assertIn("DTSTART:20261021T000000Z", text)  # even week 6
            self.assertIn("DTSTART:20260917T080000Z", text)  # custom week 1
            self.assertIn("DTSTART:20261105T080000Z", text)  # custom week 8
            self.assertIn("SUMMARY:线性代数", text)
            self.assertIn("LOCATION:二教301", text)
            for line in raw.split(b"\r\n"):
                self.assertLessEqual(len(line), 75)

    def test_invalid_course_row_is_skipped_but_valid_events_are_written(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            xlsx_path = root / "timetable.xlsx"
            ics_path = root / "timetable.ics"
            write_excel(validate_and_normalize(sample_data()), xlsx_path)
            workbook = load_workbook(xlsx_path)
            try:
                courses = workbook["Courses"]
                courses.append(["坏数据", 9, 1, 2, 1, 2, "all", "", "", "", ""])
                workbook.save(xlsx_path)
            finally:
                workbook.close()

            result = convert_excel_to_ics(xlsx_path, ics_path)
            self.assertEqual(result.event_count, 13)
            self.assertEqual(result.skipped_course_count, 1)
            self.assertIn("第 6 行：weekday = 9", result.errors[0])

    def test_zero_alarm_disables_valarm(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            xlsx_path = root / "timetable.xlsx"
            ics_path = root / "timetable.ics"
            write_excel(validate_and_normalize(sample_data()), xlsx_path)
            workbook = load_workbook(xlsx_path)
            try:
                workbook["Settings"]["B5"] = 0
                workbook.save(xlsx_path)
            finally:
                workbook.close()
            convert_excel_to_ics(xlsx_path, ics_path)
            self.assertNotIn("BEGIN:VALARM", ics_path.read_text(encoding="utf-8"))


if __name__ == "__main__":
    unittest.main()

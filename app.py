"""Friendly Windows GUI for timetable image -> Excel -> Apple Calendar.

The GUI keeps the original two-script workflow as its data layer, while making
the common path fully interactive: drop a file, review courses, set real class
times, and open the generated files. It remains a local desktop application.
"""

from __future__ import annotations

import ctypes
import json
import os
import re
import sys
import threading
import tkinter as tk
from datetime import date, datetime, time, timedelta
from pathlib import Path
from tkinter import filedialog, messagebox, simpledialog, ttk
from typing import Any
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

import customtkinter as ctk
from openpyxl import load_workbook

from excel_to_ics import EXPECTED_HEADERS, convert_excel_to_ics
from image_to_excel import (
    DEFAULT_MODEL,
    DEFAULT_NORMALIZED_JSON,
    DEFAULT_RAW_RESPONSE,
    TimetableValidationError,
    default_settings,
    request_timetable_json,
    validate_and_normalize,
    write_excel,
)

try:
    from tkinterdnd2 import DND_FILES, TkinterDnD
except ImportError:  # Selecting a file still works if optional DnD is unavailable.
    DND_FILES = None
    TkinterDnD = None


PROJECT_DIR = (
    Path(sys.executable).resolve().parent
    if getattr(sys, "frozen", False)
    else Path(__file__).resolve().parent
)
DEFAULT_OUTPUT_DIR = PROJECT_DIR / "output"
UI_SETTINGS_PATH = PROJECT_DIR / "temp" / "ui_settings.json"

COLORS = {
    "window": ("#EAF3FF", "#090F1C"),
    "surface": ("#F7FAFF", "#0F1828"),
    "sidebar": ("#F3F8FF", "#111C2E"),
    "card": ("#FFFFFF", "#162338"),
    "soft": ("#EEF6FF", "#1B2B45"),
    "border": ("#D8E7FA", "#2B3D59"),
    "text": ("#15243B", "#F5F8FF"),
    "muted": ("#67758A", "#A6B4CA"),
    "entry": ("#F8FBFF", "#0D1727"),
}
ACCENT = "#0A84FF"
ACCENT_HOVER = "#0071E3"
PURPLE = "#7C6CF2"
CYAN = "#32ADE6"
GREEN = "#34C759"
RED = "#FF453A"

COURSE_COLUMNS = (
    "course_name",
    "weekday",
    "sections",
    "weeks",
    "week_pattern",
    "custom_weeks",
    "location",
    "teacher",
    "notes",
)
APPEARANCE_TO_CTK = {"跟随系统": "System", "浅色": "Light", "深色": "Dark"}
DEFAULT_UI_SETTINGS = {
    "appearance": "浅色",
    "glass": True,
    "opacity": 96,
    "blur": 70,
}


def parse_date(value: str, label: str) -> date:
    try:
        return datetime.strptime(value.strip(), "%Y-%m-%d").date()
    except ValueError as exc:
        raise ValueError(f"{label} 应使用 YYYY-MM-DD，例如 2026-09-14") from exc


def calculate_first_week_monday(reference_monday: date, reference_week: int) -> date:
    """Calculate teaching-week 1 Monday from a known Monday and week number."""
    if reference_monday.weekday() != 0:
        raise ValueError("参考日期必须是周一")
    if reference_week < 1:
        raise ValueError("参考教学周必须大于 0")
    return reference_monday - timedelta(weeks=reference_week - 1)


def parse_time(value: str, label: str) -> time:
    text = value.strip()
    for format_string in ("%H:%M", "%H:%M:%S"):
        try:
            return datetime.strptime(text, format_string).time()
        except ValueError:
            continue
    raise ValueError(f"{label} 应使用 HH:MM，例如 08:00")


def format_date_value(value: Any) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value or "")


def format_time_value(value: Any) -> str:
    if isinstance(value, datetime):
        return value.strftime("%H:%M")
    if isinstance(value, time):
        return value.strftime("%H:%M")
    text = str(value or "")
    return text[:5] if len(text) >= 5 else text


def _load_ui_settings() -> dict[str, Any]:
    result = dict(DEFAULT_UI_SETTINGS)
    try:
        saved = json.loads(UI_SETTINGS_PATH.read_text(encoding="utf-8"))
        if isinstance(saved, dict):
            result.update({key: saved[key] for key in result if key in saved})
    except (FileNotFoundError, json.JSONDecodeError, OSError):
        pass
    if result["appearance"] not in APPEARANCE_TO_CTK:
        result["appearance"] = "浅色"
    try:
        result["opacity"] = max(80, min(100, int(float(result["opacity"]))))
    except (TypeError, ValueError):
        result["opacity"] = 96
    try:
        result["blur"] = max(0, min(100, int(float(result["blur"]))))
    except (TypeError, ValueError):
        result["blur"] = 70
    return result


def _save_ui_settings(settings: dict[str, Any]) -> None:
    UI_SETTINGS_PATH.parent.mkdir(parents=True, exist_ok=True)
    UI_SETTINGS_PATH.write_text(
        json.dumps(settings, ensure_ascii=False, indent=2), encoding="utf-8"
    )


def read_excel_for_gui(path: Path) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    """Load course rows and all Settings values for continued GUI editing."""
    workbook = load_workbook(path, data_only=True)
    try:
        if "Courses" not in workbook.sheetnames or "Settings" not in workbook.sheetnames:
            raise ValueError("Excel 必须同时包含 Courses 和 Settings")

        settings: dict[str, Any] = {}
        for key, value in workbook["Settings"].iter_rows(
            min_row=2, max_col=2, values_only=True
        ):
            key_text = str(key or "").strip()
            if key_text:
                settings[key_text] = value

        course_sheet = workbook["Courses"]
        positions: dict[str, int] = {}
        for column_index, cell in enumerate(course_sheet[1], start=1):
            header = str(cell.value or "").strip()
            if header in EXPECTED_HEADERS:
                positions[EXPECTED_HEADERS[header]] = column_index
        missing = [label for label, field in EXPECTED_HEADERS.items() if field not in positions]
        if missing:
            raise ValueError(f"Courses 缺少列：{', '.join(missing)}")

        raw_courses: list[dict[str, Any]] = []
        max_column = max(positions.values())
        for row in course_sheet.iter_rows(min_row=2, max_col=max_column, values_only=True):
            if all(value in (None, "") for value in row):
                continue

            def cell_value(field: str) -> Any:
                return row[positions[field] - 1]

            raw_courses.append(
                {
                    "course_name": cell_value("course_name"),
                    "weekday": cell_value("weekday"),
                    "start_section": cell_value("start_section"),
                    "end_section": cell_value("end_section"),
                    "start_week": cell_value("start_week"),
                    "end_week": cell_value("end_week"),
                    "week_pattern": cell_value("week_pattern"),
                    "custom_weeks": cell_value("custom_weeks"),
                    "location": cell_value("location"),
                    "teacher": cell_value("teacher"),
                    "notes": cell_value("notes"),
                }
            )
        return validate_and_normalize({"courses": raw_courses})["courses"], settings
    finally:
        workbook.close()


def validate_courses_for_calendar(
    courses: list[dict[str, Any]], available_sections: int | None = None
) -> None:
    required = {
        "weekday": "星期",
        "start_section": "开始节次",
        "end_section": "结束节次",
        "start_week": "开始周",
        "end_week": "结束周",
    }
    errors: list[str] = []
    for index, course in enumerate(courses, start=1):
        name = course.get("course_name", "")
        for field, label in required.items():
            if course.get(field) is None:
                errors.append(f"第 {index} 门课程《{name}》缺少{label}")
        if course.get("week_pattern") == "custom" and not course.get("custom_weeks"):
            errors.append(f"第 {index} 门课程《{name}》缺少指定周")
        if (
            available_sections is not None
            and course.get("end_section") is not None
            and course["end_section"] > available_sections
        ):
            errors.append(
                f"第 {index} 门课程《{name}》使用第 {course['end_section']} 节，"
                f"但当前只设置了 {available_sections} 节"
            )
    if errors:
        raise ValueError("\n".join(errors))


def apply_course_patch(
    courses: list[dict[str, Any]], indices: list[int], patch: dict[str, Any]
) -> list[dict[str, Any]]:
    """Return courses with a validated patch applied to selected rows."""
    updated = [dict(course) for course in courses]
    for index in indices:
        raw = dict(updated[index])
        raw.update(patch)
        normalized = validate_and_normalize({"courses": [raw]})["courses"][0]
        if normalized["week_pattern"] == "custom" and any(
            week < normalized["start_week"] or week > normalized["end_week"]
            for week in normalized["custom_weeks"]
        ):
            raise ValueError("指定周必须位于开始周和结束周之间")
        updated[index] = normalized
    return updated


def _windows_backdrop(
    window: tk.Misc, enabled: bool, opacity: int, blur: int
) -> None:
    """Apply a readable Windows 11 backdrop plus independent window opacity."""
    try:
        window.update_idletasks()
        hwnd = ctypes.windll.user32.GetParent(window.winfo_id()) or window.winfo_id()
        window.configure(fg_color=COLORS["window"])
        if not enabled or blur <= 0:
            backdrop_type = 1  # None
        elif blur < 35:
            backdrop_type = 2  # Mica: restrained
        elif blur < 70:
            backdrop_type = 4  # Mica Alt: medium
        else:
            backdrop_type = 3  # Desktop Acrylic: strongest blur
        backdrop = ctypes.c_int(backdrop_type)
        ctypes.windll.dwmapi.DwmSetWindowAttribute(
            hwnd, 38, ctypes.byref(backdrop), ctypes.sizeof(backdrop)
        )
        dark = ctypes.c_int(1 if ctk.get_appearance_mode() == "Dark" else 0)
        ctypes.windll.dwmapi.DwmSetWindowAttribute(
            hwnd, 20, ctypes.byref(dark), ctypes.sizeof(dark)
        )
    except (AttributeError, OSError, ctypes.ArgumentError):
        pass
    try:
        window.attributes("-alpha", opacity / 100 if enabled else 1.0)
    except tk.TclError:
        pass


if TkinterDnD is not None:

    class AppRoot(ctk.CTk, TkinterDnD.DnDWrapper):  # type: ignore[misc]
        def __init__(self) -> None:
            ctk.CTk.__init__(self)
            self.TkdndVersion = TkinterDnD._require(self)

else:

    class AppRoot(ctk.CTk):
        pass


class AppButton(ctk.CTkButton):
    """Consistent application button without motion or resizing effects."""

    def __init__(
        self,
        master: Any,
        *,
        width: int = 178,
        height: int = 44,
        font_size: int = 14,
        **kwargs: Any,
    ) -> None:
        kwargs.setdefault(
            "font", ctk.CTkFont(family="Segoe UI", size=font_size, weight="bold")
        )
        kwargs.setdefault("corner_radius", 14)
        super().__init__(master, width=width, height=height, **kwargs)


class CourseDialog:
    FIELD_SPECS = (
        ("course_name", "课程名称"),
        ("weekday", "星期（1-7）"),
        ("start_section", "开始节次"),
        ("end_section", "结束节次"),
        ("start_week", "开始周"),
        ("end_week", "结束周"),
        ("week_pattern", "周期"),
        ("custom_weeks", "指定周"),
        ("location", "地点"),
        ("teacher", "教师"),
        ("notes", "备注"),
    )

    def __init__(self, parent: tk.Misc, course: dict[str, Any] | None = None):
        self.result: dict[str, Any] | None = None
        self.window = ctk.CTkToplevel(parent)
        self.window.title("编辑课程" if course else "添加课程")
        self.window.geometry("540x700")
        self.window.minsize(500, 590)
        self.window.transient(parent)
        self.window.grab_set()
        self.window.grid_columnconfigure(0, weight=1)
        self.window.grid_rowconfigure(1, weight=1)
        self.variables: dict[str, tk.StringVar] = {}
        ctk.CTkLabel(
            self.window,
            text="编辑课程" if course else "添加一门课程",
            text_color=COLORS["text"],
            font=ctk.CTkFont("Segoe UI", 22, "bold"),
        ).grid(row=0, column=0, sticky="w", padx=26, pady=(22, 10))
        form = ctk.CTkScrollableFrame(self.window, fg_color="transparent", corner_radius=0)
        form.grid(row=1, column=0, sticky="nsew", padx=22)
        form.grid_columnconfigure(1, weight=1)
        course = course or {}
        for row_number, (field, label) in enumerate(self.FIELD_SPECS):
            ctk.CTkLabel(
                form, text=label, anchor="w", text_color=COLORS["muted"]
            ).grid(row=row_number, column=0, sticky="w", padx=(4, 16), pady=7)
            value = course.get(field, "")
            if field == "custom_weeks" and isinstance(value, list):
                value = ",".join(map(str, value))
            variable = tk.StringVar(value=str(value if value is not None else ""))
            self.variables[field] = variable
            if field == "week_pattern":
                if not variable.get():
                    variable.set("all")
                widget: ctk.CTkBaseClass = ctk.CTkOptionMenu(
                    form,
                    variable=variable,
                    values=["all", "odd", "even", "custom"],
                    fg_color=COLORS["soft"],
                    button_color=ACCENT,
                    text_color=COLORS["text"],
                )
            else:
                widget = ctk.CTkEntry(
                    form,
                    textvariable=variable,
                    height=38,
                    corner_radius=12,
                    fg_color=COLORS["entry"],
                    border_color=COLORS["border"],
                    text_color=COLORS["text"],
                )
            widget.grid(row=row_number, column=1, sticky="ew", pady=7)
        ctk.CTkLabel(
            form,
            text="custom 时填写英文逗号分隔周次，例如 1,3,5,8",
            anchor="w",
            text_color=COLORS["muted"],
            font=ctk.CTkFont("Segoe UI", 12),
        ).grid(row=len(self.FIELD_SPECS), column=0, columnspan=2, sticky="w", pady=(4, 12))
        actions = ctk.CTkFrame(self.window, fg_color="transparent")
        actions.grid(row=2, column=0, sticky="e", padx=24, pady=20)
        ctk.CTkButton(
            actions,
            text="取消",
            width=96,
            fg_color=COLORS["soft"],
            hover_color=COLORS["border"],
            text_color=COLORS["text"],
            command=self.window.destroy,
        ).pack(side="left", padx=6)
        ctk.CTkButton(
            actions,
            text="保存",
            width=112,
            fg_color=ACCENT,
            hover_color=ACCENT_HOVER,
            command=self._save,
        ).pack(side="left", padx=6)
        self.window.bind("<Escape>", lambda _event: self.window.destroy())
        self.window.bind("<Control-Return>", lambda _event: self._save())
        self.window.after(40, lambda: self._center(parent))
        self.window.wait_window()

    def _center(self, parent: tk.Misc) -> None:
        self.window.update_idletasks()
        x = parent.winfo_rootx() + max(0, (parent.winfo_width() - self.window.winfo_width()) // 2)
        y = parent.winfo_rooty() + max(0, (parent.winfo_height() - self.window.winfo_height()) // 2)
        self.window.geometry(f"+{x}+{y}")

    def _save(self) -> None:
        raw = {field: variable.get() for field, variable in self.variables.items()}
        try:
            normalized = validate_and_normalize({"courses": [raw]})["courses"][0]
            validate_courses_for_calendar([normalized])
            if normalized["week_pattern"] == "custom" and any(
                week < normalized["start_week"] or week > normalized["end_week"]
                for week in normalized["custom_weeks"]
            ):
                raise ValueError("指定周必须位于开始周和结束周之间")
            self.result = normalized
            self.window.destroy()
        except (TimetableValidationError, TypeError, ValueError) as exc:
            messagebox.showerror("课程数据有误", str(exc), parent=self.window)


class BatchCourseDialog:
    FIELD_SPECS = (
        ("weekday", "星期（1-7）"),
        ("start_section", "开始节次"),
        ("end_section", "结束节次"),
        ("start_week", "开始周"),
        ("end_week", "结束周"),
        ("week_pattern", "周期"),
        ("custom_weeks", "指定周"),
        ("location", "地点"),
        ("teacher", "教师"),
        ("notes", "备注"),
    )

    def __init__(self, parent: tk.Misc, count: int):
        self.result: dict[str, Any] | None = None
        self.window = ctk.CTkToplevel(parent)
        self.window.title("批量编辑课程")
        self.window.geometry("630x680")
        self.window.minsize(560, 560)
        self.window.transient(parent)
        self.window.grab_set()
        self.window.grid_columnconfigure(0, weight=1)
        self.window.grid_rowconfigure(2, weight=1)
        self.enabled: dict[str, tk.BooleanVar] = {}
        self.values: dict[str, tk.StringVar] = {}
        ctk.CTkLabel(
            self.window,
            text=f"批量编辑 {count} 门课程",
            text_color=COLORS["text"],
            font=ctk.CTkFont("Segoe UI", 22, "bold"),
        ).grid(row=0, column=0, sticky="w", padx=28, pady=(24, 4))
        ctk.CTkLabel(
            self.window,
            text="勾选要统一修改的字段；勾选后留空可清空地点、教师或备注。",
            text_color=COLORS["muted"],
        ).grid(row=1, column=0, sticky="w", padx=28, pady=(0, 14))
        form = ctk.CTkScrollableFrame(self.window, fg_color="transparent")
        form.grid(row=2, column=0, sticky="nsew", padx=24)
        form.grid_columnconfigure(2, weight=1)
        for row_number, (field, label) in enumerate(self.FIELD_SPECS):
            enabled = tk.BooleanVar(value=False)
            value = tk.StringVar(value="all" if field == "week_pattern" else "")
            self.enabled[field] = enabled
            self.values[field] = value
            ctk.CTkCheckBox(
                form,
                text="",
                variable=enabled,
                width=26,
                checkbox_width=22,
                checkbox_height=22,
                fg_color=ACCENT,
            ).grid(row=row_number, column=0, padx=(2, 8), pady=8)
            ctk.CTkLabel(form, text=label, width=110, anchor="w").grid(
                row=row_number, column=1, sticky="w", padx=(0, 12), pady=8
            )
            if field == "week_pattern":
                widget: ctk.CTkBaseClass = ctk.CTkOptionMenu(
                    form,
                    variable=value,
                    values=["all", "odd", "even", "custom"],
                    fg_color=COLORS["soft"],
                    button_color=ACCENT,
                    text_color=COLORS["text"],
                )
            else:
                widget = ctk.CTkEntry(
                    form,
                    textvariable=value,
                    height=38,
                    fg_color=COLORS["entry"],
                    border_color=COLORS["border"],
                    text_color=COLORS["text"],
                )
            widget.grid(row=row_number, column=2, sticky="ew", pady=8)
        actions = ctk.CTkFrame(self.window, fg_color="transparent")
        actions.grid(row=3, column=0, sticky="e", padx=26, pady=20)
        ctk.CTkButton(
            actions,
            text="取消",
            width=96,
            fg_color=COLORS["soft"],
            hover_color=COLORS["border"],
            text_color=COLORS["text"],
            command=self.window.destroy,
        ).pack(side="left", padx=6)
        ctk.CTkButton(
            actions,
            text="应用修改",
            width=124,
            fg_color=ACCENT,
            hover_color=ACCENT_HOVER,
            command=self._save,
        ).pack(side="left", padx=6)
        self.window.bind("<Escape>", lambda _event: self.window.destroy())
        self.window.wait_window()

    def _save(self) -> None:
        patch = {
            field: self.values[field].get()
            for field in self.values
            if self.enabled[field].get()
        }
        if not patch:
            messagebox.showinfo("尚未选择字段", "请至少勾选一个需要修改的字段。", parent=self.window)
            return
        self.result = patch
        self.window.destroy()


class TimetableApp:
    PAGE_TITLES = {
        "import": "图片识别",
        "courses": "课程编辑",
        "schedule": "学期与节次",
        "export": "导出结果",
        "settings": "外观设置",
    }

    def __init__(self, root: AppRoot):
        self.root = root
        self.ui_settings = _load_ui_settings()
        ctk.set_appearance_mode(APPEARANCE_TO_CTK[self.ui_settings["appearance"]])
        self.root.title("课表日历助手")
        self._configure_window_size()
        self.root.configure(fg_color=COLORS["window"])
        self.courses: list[dict[str, Any]] = []
        self.busy = False
        self.in_workspace = False
        self.current_page = "import"

        self.image_path = tk.StringVar()
        self.api_key = tk.StringVar(value=os.getenv("OPENAI_API_KEY", ""))
        self.model = tk.StringVar(value=os.getenv("OPENAI_MODEL", DEFAULT_MODEL))
        self.reference_date = tk.StringVar(value="2026-09-14")
        self.reference_week = tk.StringVar(value="1")
        self.first_week_monday = tk.StringVar(value="2026-09-14")
        self.first_week_hint = tk.StringVar(value="")
        self.semester_name = tk.StringVar(value="2026秋季学期")
        self.semester_end = tk.StringVar(value="2027-01-15")
        self.alarm_minutes = tk.StringVar(value="15")
        self.timezone = tk.StringVar(value="Asia/Shanghai")
        self.output_dir = tk.StringVar(value=str(DEFAULT_OUTPUT_DIR))
        self.course_count = tk.StringVar(value="0 门课程")
        self.status = tk.StringVar(value="点击主页的“开始”进入课表处理流程。")
        self.completion_text = tk.StringVar(value="还没有生成文件")
        self.appearance_var = tk.StringVar(value=self.ui_settings["appearance"])
        self.glass_var = tk.BooleanVar(value=bool(self.ui_settings["glass"]))
        self.opacity_var = tk.DoubleVar(value=float(self.ui_settings["opacity"]))
        self.blur_var = tk.DoubleVar(value=float(self.ui_settings["blur"]))
        self.opacity_label = tk.StringVar(value=f"{int(self.opacity_var.get())}%")
        self.blur_label = tk.StringVar(value=f"{int(self.blur_var.get())}%")
        self.section_variables: dict[int, tuple[tk.BooleanVar, tk.StringVar, tk.StringVar]] = {}
        self.current_excel_path = DEFAULT_OUTPUT_DIR / "timetable.xlsx"
        self.current_ics_path = DEFAULT_OUTPUT_DIR / "timetable.ics"

        self._configure_tree_style()
        self._build_shell()
        self._bind_shortcuts()
        self._load_default_section_times()
        self._update_reference_hint()
        self.root.protocol("WM_DELETE_WINDOW", self._close)
        self.root.after(180, self._apply_window_effects)

    def _configure_window_size(self) -> None:
        screen_width = self.root.winfo_screenwidth()
        screen_height = self.root.winfo_screenheight()
        width = min(1280, max(900, screen_width - 100))
        height = min(840, max(660, screen_height - 100))
        min_width = min(1040, max(860, screen_width - 140))
        min_height = min(700, max(620, screen_height - 140))
        x = max(0, (screen_width - width) // 2)
        y = max(0, (screen_height - height) // 2)
        self.root.geometry(f"{width}x{height}+{x}+{y}")
        self.root.minsize(min_width, min_height)

    def _bind_shortcuts(self) -> None:
        pages = {
            "1": "import",
            "2": "courses",
            "3": "schedule",
            "4": "export",
            "comma": "settings",
        }
        for key, page in pages.items():
            self.root.bind(
                f"<Control-Key-{key}>",
                lambda _event, target=page: self.enter_workspace(target),
                add="+",
            )
        self.root.bind(
            "<Escape>",
            lambda _event: self.exit_workspace() if self.in_workspace else None,
            add="+",
        )

    def _build_shell(self) -> None:
        self.viewport = ctk.CTkFrame(self.root, fg_color="transparent", corner_radius=0)
        self.viewport.pack(fill="both", expand=True)
        self._build_home()
        self._build_workspace()
        self.home.place(relx=0, rely=0, relwidth=1, relheight=1)

    def _build_home(self) -> None:
        self.home = ctk.CTkFrame(self.viewport, fg_color="transparent", corner_radius=0)
        top = ctk.CTkFrame(self.home, fg_color="transparent")
        top.pack(fill="x", padx=48, pady=(34, 14))
        ctk.CTkLabel(
            top,
            text="课表日历助手",
            text_color=COLORS["text"],
            font=ctk.CTkFont("Segoe UI", 30, "bold"),
        ).pack(side="left")
        AppButton(
            top,
            text="⚙  外观",
            width=112,
            height=42,
            font_size=13,
            fg_color=COLORS["card"],
            hover_color=COLORS["soft"],
            border_width=1,
            border_color=COLORS["border"],
            text_color=COLORS["text"],
            command=lambda: self.enter_workspace("settings"),
        ).pack(side="right")
        hero = ctk.CTkFrame(
            self.home,
            fg_color=("#DCEEFF", "#14243C"),
            corner_radius=36,
            border_width=1,
            border_color=COLORS["border"],
        )
        hero.pack(fill="both", expand=True, padx=48, pady=(8, 42))
        hero.grid_columnconfigure(0, weight=1)
        hero.grid_rowconfigure(0, weight=1)
        hero.grid_rowconfigure(2, weight=1)
        content = ctk.CTkFrame(hero, fg_color="transparent")
        content.grid(row=1, column=0, padx=44, pady=34)
        ctk.CTkLabel(
            content,
            text="为真正可用的课程日历而做",
            width=224,
            height=34,
            corner_radius=17,
            fg_color=("#FFFFFF", "#203653"),
            text_color=ACCENT,
            font=ctk.CTkFont("Segoe UI", 12, "bold"),
        ).pack(pady=(0, 18))
        ctk.CTkLabel(
            content,
            text="让课表图片里的信息，\n变成你确认过的日程。",
            justify="center",
            text_color=COLORS["text"],
            font=ctk.CTkFont("Segoe UI", 35, "bold"),
        ).pack()
        ctk.CTkLabel(
            content,
            text=(
                "课表图片看起来清楚，真正导入日历时却常在教学周和节次时间上出错。\n"
                "这个工具的初衷，是在生成文件前，把每个关键字段都交给你检查和修改。"
            ),
            justify="center",
            wraplength=620,
            text_color=COLORS["muted"],
            font=ctk.CTkFont("Segoe UI", 14),
        ).pack(pady=(16, 24))
        flow = ctk.CTkFrame(
            content,
            fg_color=("#F6FAFF", "#1A2D49"),
            corner_radius=18,
            border_width=1,
            border_color=COLORS["border"],
        )
        flow.pack(fill="x", pady=(0, 28))
        ctk.CTkLabel(
            flow,
            text="拖入图片   →   检查课程   →   设置教学周与节次   →   生成 Excel 和 ICS",
            justify="center",
            wraplength=620,
            text_color=COLORS["text"],
            font=ctk.CTkFont("Segoe UI", 13, "bold"),
        ).pack(fill="x", padx=24, pady=17)
        AppButton(
            content,
            text="开始",
            width=320,
            height=62,
            font_size=18,
            fg_color=ACCENT,
            hover_color=ACCENT_HOVER,
            corner_radius=20,
            command=lambda: self.enter_workspace("import"),
        ).pack()
        ctk.CTkLabel(
            content,
            text="本地桌面工具 · 识别后先核对 · 输出文件保存在本机",
            text_color=COLORS["muted"],
            font=ctk.CTkFont("Segoe UI", 11),
        ).pack(pady=(16, 0))

    def _build_workspace(self) -> None:
        self.workspace = ctk.CTkFrame(
            self.viewport, fg_color=COLORS["surface"], corner_radius=0
        )
        self.workspace.grid_columnconfigure(1, weight=1)
        self.workspace.grid_rowconfigure(0, weight=1)
        sidebar = ctk.CTkFrame(
            self.workspace,
            width=218,
            fg_color=COLORS["sidebar"],
            corner_radius=0,
            border_width=1,
            border_color=COLORS["border"],
        )
        sidebar.grid(row=0, column=0, sticky="nsew")
        sidebar.grid_propagate(False)
        ctk.CTkLabel(
            sidebar,
            text="课表助手",
            text_color=COLORS["text"],
            font=ctk.CTkFont("Segoe UI", 21, "bold"),
        ).pack(anchor="w", padx=22, pady=(26, 24))
        nav_specs = (
            ("⌂  主页", "home"),
            ("◫  图片识别", "import"),
            ("≡  课程编辑", "courses"),
            ("◷  学期与节次", "schedule"),
            ("⇧  导出结果", "export"),
            ("⚙  外观设置", "settings"),
        )
        self.nav_buttons: dict[str, AppButton] = {}
        for label, page in nav_specs:
            command = self.exit_workspace if page == "home" else lambda target=page: self.navigate(target)
            button = AppButton(
                sidebar,
                text=label,
                anchor="w",
                width=178,
                height=45,
                font_size=13,
                fg_color="transparent",
                hover_color=COLORS["soft"],
                text_color=COLORS["text"],
                command=command,
            )
            button.pack(padx=18, pady=4)
            self.nav_buttons[page] = button
        ctk.CTkLabel(
            sidebar,
            text="Excel 仍是可编辑的数据源\nICS 由当前设置生成",
            justify="left",
            text_color=COLORS["muted"],
            font=ctk.CTkFont("Segoe UI", 11),
        ).pack(side="bottom", anchor="w", padx=22, pady=24)
        content = ctk.CTkFrame(self.workspace, fg_color="transparent", corner_radius=0)
        content.grid(row=0, column=1, sticky="nsew")
        content.grid_columnconfigure(0, weight=1)
        content.grid_rowconfigure(1, weight=1)
        topbar = ctk.CTkFrame(content, fg_color="transparent", height=76)
        topbar.grid(row=0, column=0, sticky="ew", padx=26, pady=(16, 4))
        self.page_title = ctk.CTkLabel(
            topbar,
            text="图片识别",
            text_color=COLORS["text"],
            font=ctk.CTkFont("Segoe UI", 26, "bold"),
        )
        self.page_title.pack(side="left")
        ctk.CTkLabel(
            topbar,
            text="V2.1 本地版",
            width=86,
            height=30,
            corner_radius=15,
            fg_color=COLORS["soft"],
            text_color=ACCENT,
            font=ctk.CTkFont("Segoe UI", 11, "bold"),
        ).pack(side="right")
        self.page_host = ctk.CTkFrame(content, fg_color="transparent")
        self.page_host.grid(row=1, column=0, sticky="nsew", padx=24, pady=6)
        self.page_host.grid_columnconfigure(0, weight=1)
        self.page_host.grid_rowconfigure(0, weight=1)
        self.pages: dict[str, ctk.CTkFrame] = {}
        self.page_containers: dict[str, ctk.CTkFrame] = {}
        self._build_import_page()
        self._build_courses_page()
        self._build_schedule_page()
        self._build_export_page()
        self._build_settings_page()
        footer = ctk.CTkFrame(
            content,
            height=52,
            fg_color=COLORS["card"],
            corner_radius=18,
            border_width=1,
            border_color=COLORS["border"],
        )
        footer.grid(row=2, column=0, sticky="ew", padx=24, pady=(4, 18))
        footer.grid_columnconfigure(0, weight=1)
        ctk.CTkLabel(
            footer,
            textvariable=self.status,
            anchor="w",
            text_color=COLORS["muted"],
            font=ctk.CTkFont("Segoe UI", 12),
        ).grid(row=0, column=0, sticky="ew", padx=18, pady=13)
        self.progress = ctk.CTkProgressBar(
            footer, width=150, height=8, progress_color=ACCENT, mode="indeterminate"
        )
        self.progress.grid(row=0, column=1, padx=18)
        self.progress.set(0)

    def _new_page(self, name: str, scrollable: bool = False) -> ctk.CTkFrame:
        # A fixed wrapper gives every page one reliable sibling for show/hide.
        # CTkScrollableFrame itself owns an extra parent frame and canvas, so
        # raising the inner widget alone can leave a later page visible.
        container = ctk.CTkFrame(
            self.page_host, fg_color="transparent", corner_radius=0
        )
        container.grid(row=0, column=0, sticky="nsew")
        if scrollable:
            page: ctk.CTkFrame = ctk.CTkScrollableFrame(
                container, fg_color="transparent", corner_radius=0
            )
            page.pack(fill="both", expand=True)
        else:
            page = container
        self.page_containers[name] = container
        self.pages[name] = page
        return page

    def _card(self, parent: Any, *, corner_radius: int = 22) -> ctk.CTkFrame:
        return ctk.CTkFrame(
            parent,
            fg_color=COLORS["card"],
            corner_radius=corner_radius,
            border_width=1,
            border_color=COLORS["border"],
        )

    def _section_title(self, parent: Any, title: str, subtitle: str = "") -> None:
        ctk.CTkLabel(
            parent,
            text=title,
            anchor="w",
            text_color=COLORS["text"],
            font=ctk.CTkFont("Segoe UI", 20, "bold"),
        ).pack(fill="x", padx=4, pady=(2, 2))
        if subtitle:
            ctk.CTkLabel(
                parent,
                text=subtitle,
                anchor="w",
                justify="left",
                text_color=COLORS["muted"],
                font=ctk.CTkFont("Segoe UI", 13),
            ).pack(fill="x", padx=4, pady=(0, 16))

    def _entry(self, parent: Any, variable: tk.Variable, *, show: str | None = None) -> ctk.CTkEntry:
        return ctk.CTkEntry(
            parent,
            textvariable=variable,
            show=show or "",
            height=42,
            corner_radius=13,
            fg_color=COLORS["entry"],
            border_color=COLORS["border"],
            text_color=COLORS["text"],
        )

    def _build_import_page(self) -> None:
        page = self._new_page("import", scrollable=True)
        self._section_title(
            page,
            "选择课表来源",
            "拖入一张清晰的课表图开始识别，或者加载之前保存的 Excel 继续修改。",
        )
        body = ctk.CTkFrame(page, fg_color="transparent")
        body.pack(fill="both", expand=True)
        body.grid_columnconfigure(0, weight=1, uniform="import")
        body.grid_columnconfigure(1, weight=1, uniform="import")
        self.import_body = body
        drop_card = self._card(body)
        drop_card.grid(row=0, column=0, sticky="nsew", padx=(0, 8), pady=4)
        self.import_drop_card = drop_card
        drop_card.grid_columnconfigure(0, weight=1)
        drop_card.grid_rowconfigure(1, weight=1)
        ctk.CTkLabel(
            drop_card,
            text="课表图片",
            anchor="w",
            text_color=COLORS["text"],
            font=ctk.CTkFont("Segoe UI", 16, "bold"),
        ).grid(row=0, column=0, sticky="ew", padx=22, pady=(20, 12))
        self.drop_zone = ctk.CTkFrame(
            drop_card,
            height=255,
            fg_color=COLORS["soft"],
            corner_radius=22,
            border_width=2,
            border_color=("#9CCBFF", "#315B89"),
        )
        self.drop_zone.grid(row=1, column=0, sticky="nsew", padx=22)
        self.drop_zone.grid_columnconfigure(0, weight=1)
        self.drop_zone.grid_rowconfigure(0, weight=1)
        drop_content = ctk.CTkFrame(self.drop_zone, fg_color="transparent")
        drop_content.grid(row=0, column=0)
        self.drop_icon = ctk.CTkLabel(
            drop_content,
            text="⇩",
            text_color=ACCENT,
            font=ctk.CTkFont("Segoe UI", 38, "bold"),
        )
        self.drop_icon.pack()
        self.drop_title = ctk.CTkLabel(
            drop_content,
            text="请拖入课表图片",
            text_color=COLORS["text"],
            font=ctk.CTkFont("Segoe UI", 17, "bold"),
        )
        self.drop_title.pack(pady=(4, 4))
        self.drop_detail = ctk.CTkLabel(
            drop_content,
            text="支持 PNG、JPEG、WEBP、GIF\n也可以点击这里选择文件",
            justify="center",
            text_color=COLORS["muted"],
            font=ctk.CTkFont("Segoe UI", 12),
        )
        self.drop_detail.pack()
        for widget in (self.drop_zone, drop_content, self.drop_icon, self.drop_title, self.drop_detail):
            widget.bind("<Button-1>", lambda _event: self.select_image())
        self._enable_drop_target()
        ctk.CTkLabel(
            drop_card,
            textvariable=self.image_path,
            anchor="w",
            text_color=COLORS["muted"],
            font=ctk.CTkFont("Segoe UI", 11),
            wraplength=430,
        ).grid(row=2, column=0, sticky="ew", padx=24, pady=(12, 4))
        image_actions = ctk.CTkFrame(drop_card, fg_color="transparent")
        image_actions.grid(row=3, column=0, sticky="ew", padx=22, pady=(8, 22))
        ctk.CTkButton(
            image_actions,
            text="选择图片",
            fg_color=ACCENT,
            hover_color=ACCENT_HOVER,
            command=self.select_image,
        ).pack(side="left")
        ctk.CTkButton(
            image_actions,
            text="加载已有 Excel",
            fg_color=COLORS["soft"],
            hover_color=COLORS["border"],
            text_color=COLORS["text"],
            command=self.load_existing_excel,
        ).pack(side="left", padx=10)
        api_card = self._card(body)
        api_card.grid(row=0, column=1, sticky="nsew", padx=(8, 0), pady=4)
        self.import_api_card = api_card
        api_card.grid_columnconfigure(0, weight=1)
        ctk.CTkLabel(
            api_card,
            text="识别设置",
            anchor="w",
            text_color=COLORS["text"],
            font=ctk.CTkFont("Segoe UI", 16, "bold"),
        ).grid(row=0, column=0, sticky="ew", padx=22, pady=(20, 16))
        ctk.CTkLabel(api_card, text="OpenAI API Key", anchor="w").grid(
            row=1, column=0, sticky="ew", padx=22
        )
        self._entry(api_card, self.api_key, show="•").grid(
            row=2, column=0, sticky="ew", padx=22, pady=(6, 5)
        )
        ctk.CTkLabel(
            api_card,
            text="只在本次运行中使用，不会写入配置文件。",
            anchor="w",
            text_color=COLORS["muted"],
            font=ctk.CTkFont("Segoe UI", 11),
        ).grid(row=3, column=0, sticky="ew", padx=22)
        ctk.CTkLabel(api_card, text="视觉模型", anchor="w").grid(
            row=4, column=0, sticky="ew", padx=22, pady=(18, 0)
        )
        self._entry(api_card, self.model).grid(
            row=5, column=0, sticky="ew", padx=22, pady=(6, 14)
        )
        ctk.CTkLabel(
            api_card,
            text="识别完成后不会立刻生成日历。你会先进入课程编辑页逐项确认。",
            anchor="w",
            justify="left",
            wraplength=420,
            text_color=COLORS["muted"],
            font=ctk.CTkFont("Segoe UI", 12),
        ).grid(row=6, column=0, sticky="ew", padx=22, pady=(4, 18))
        self.recognize_button = AppButton(
            api_card,
            text="开始识别  →",
            width=210,
            height=48,
            fg_color=ACCENT,
            hover_color=ACCENT_HOVER,
            command=self.recognize_image,
        )
        self.recognize_button.grid(row=7, column=0, sticky="e", padx=22, pady=(0, 22))
        body.bind("<Configure>", self._layout_import_cards)

    def _layout_import_cards(self, event: tk.Event) -> None:
        """Stack the two import cards before either card becomes cramped."""
        if event.width < 760:
            self.import_body.grid_columnconfigure(1, weight=0, uniform="")
            self.import_drop_card.grid_configure(
                row=0, column=0, sticky="nsew", padx=0, pady=(4, 8)
            )
            self.import_api_card.grid_configure(
                row=1, column=0, sticky="nsew", padx=0, pady=(8, 4)
            )
        else:
            self.import_body.grid_columnconfigure(1, weight=1, uniform="import")
            self.import_drop_card.grid_configure(
                row=0, column=0, sticky="nsew", padx=(0, 8), pady=4
            )
            self.import_api_card.grid_configure(
                row=0, column=1, sticky="nsew", padx=(8, 0), pady=4
            )

    def _enable_drop_target(self) -> None:
        if not DND_FILES:
            return
        try:
            self.drop_zone.drop_target_register(DND_FILES)  # type: ignore[attr-defined]
            self.drop_zone.dnd_bind("<<Drop>>", self._on_drop)  # type: ignore[attr-defined]
        except (AttributeError, tk.TclError):
            self.drop_detail.configure(text="点击这里选择图片（当前环境不支持拖放）")

    def _build_courses_page(self) -> None:
        page = self._new_page("courses")
        page.grid_columnconfigure(0, weight=1)
        page.grid_rowconfigure(2, weight=1)
        ctk.CTkLabel(
            page,
            text="检查和编辑课程",
            anchor="w",
            text_color=COLORS["text"],
            font=ctk.CTkFont("Segoe UI", 20, "bold"),
        ).grid(row=0, column=0, sticky="ew", padx=4)
        actions = ctk.CTkFrame(page, fg_color="transparent")
        actions.grid(row=1, column=0, sticky="ew", pady=(8, 12))
        summary_row = ctk.CTkFrame(actions, fg_color="transparent")
        summary_row.pack(fill="x")
        ctk.CTkLabel(
            summary_row,
            textvariable=self.course_count,
            text_color=ACCENT,
            font=ctk.CTkFont("Segoe UI", 13, "bold"),
        ).pack(side="left", padx=(4, 12))
        ctk.CTkLabel(
            summary_row,
            text="Ctrl/Shift 可多选，双击编辑单行",
            text_color=COLORS["muted"],
            font=ctk.CTkFont("Segoe UI", 12),
        ).pack(side="left")
        button_row = ctk.CTkFrame(actions, fg_color="transparent")
        button_row.pack(fill="x", pady=(8, 0))
        action_specs = (
            ("删除所选", self.delete_selected_courses, RED, "#FFFFFF"),
            ("批量编辑", self.batch_edit_courses, PURPLE, "#FFFFFF"),
            ("编辑所选", self.edit_selected_course, COLORS["soft"], COLORS["text"]),
            ("＋ 添加课程", self.add_course, ACCENT, "#FFFFFF"),
        )
        for text, command, color, text_color in action_specs:
            ctk.CTkButton(
                button_row,
                text=text,
                width=104,
                height=36,
                corner_radius=12,
                fg_color=color,
                hover_color=ACCENT_HOVER if color == ACCENT else COLORS["border"],
                text_color=text_color,
                command=command,
            ).pack(side="right", padx=4)
        tree_card = self._card(page, corner_radius=20)
        tree_card.grid(row=2, column=0, sticky="nsew")
        tree_card.grid_columnconfigure(0, weight=1)
        tree_card.grid_rowconfigure(0, weight=1)
        self.course_tree = ttk.Treeview(
            tree_card,
            columns=COURSE_COLUMNS,
            show="headings",
            selectmode="extended",
            style="Timetable.Treeview",
        )
        headings = {
            "course_name": "课程名称",
            "weekday": "星期",
            "sections": "节次",
            "weeks": "周次",
            "week_pattern": "周期",
            "custom_weeks": "指定周",
            "location": "地点",
            "teacher": "教师",
            "notes": "备注",
        }
        widths = {
            "course_name": 160,
            "weekday": 58,
            "sections": 70,
            "weeks": 72,
            "week_pattern": 74,
            "custom_weeks": 105,
            "location": 135,
            "teacher": 92,
            "notes": 160,
        }
        for column in COURSE_COLUMNS:
            self.course_tree.heading(column, text=headings[column])
            self.course_tree.column(column, width=widths[column], minwidth=52, stretch=True)
        vertical = ttk.Scrollbar(tree_card, orient="vertical", command=self.course_tree.yview)
        horizontal = ttk.Scrollbar(tree_card, orient="horizontal", command=self.course_tree.xview)
        self.course_tree.configure(yscrollcommand=vertical.set, xscrollcommand=horizontal.set)
        self.course_tree.grid(row=0, column=0, sticky="nsew", padx=(12, 0), pady=(12, 0))
        vertical.grid(row=0, column=1, sticky="ns", padx=(0, 12), pady=(12, 0))
        horizontal.grid(row=1, column=0, sticky="ew", padx=(12, 0), pady=(0, 12))
        self.course_tree.bind("<Double-1>", lambda _event: self.edit_selected_course())
        bottom = ctk.CTkFrame(page, fg_color="transparent")
        bottom.grid(row=3, column=0, sticky="ew", pady=(12, 0))
        ctk.CTkButton(
            bottom,
            text="下一步：设置日期和节次  →",
            width=230,
            height=42,
            fg_color=ACCENT,
            hover_color=ACCENT_HOVER,
            command=lambda: self.navigate("schedule"),
        ).pack(side="right")

    def _build_schedule_page(self) -> None:
        page = self._new_page("schedule", scrollable=True)
        self._section_title(
            page,
            "设置学期和真实时间",
            "第一教学周周一可以直接编辑，也可以用一个已知周自动反算。节次数量默认 12，可自由添加或删除。",
        )
        semester_card = self._card(page)
        semester_card.pack(fill="x", pady=(0, 14))
        for column in (1, 3):
            semester_card.grid_columnconfigure(column, weight=1)
        fields = (
            ("学期名称", self.semester_name, 0, 0),
            ("学期结束日期", self.semester_end, 0, 2),
            ("第一教学周周一", self.first_week_monday, 1, 0),
            ("课前提醒（分钟）", self.alarm_minutes, 1, 2),
            ("时区", self.timezone, 2, 0),
        )
        for label, variable, row, column in fields:
            ctk.CTkLabel(
                semester_card, text=label, anchor="w", text_color=COLORS["muted"]
            ).grid(row=row * 2, column=column, sticky="w", padx=(22, 12), pady=(18 if row == 0 else 8, 3))
            self._entry(semester_card, variable).grid(
                row=row * 2 + 1,
                column=column,
                columnspan=2 if label == "时区" else 1,
                sticky="ew",
                padx=(22, 18) if column == 0 else (0, 22),
                pady=(0, 10),
            )
        helper = ctk.CTkFrame(semester_card, fg_color=COLORS["soft"], corner_radius=16)
        helper.grid(row=6, column=0, columnspan=4, sticky="ew", padx=22, pady=(8, 22))
        helper.grid_columnconfigure(1, weight=1)
        helper.grid_columnconfigure(3, weight=1)
        ctk.CTkLabel(helper, text="已知某周的周一", text_color=COLORS["muted"]).grid(
            row=0, column=0, sticky="w", padx=(16, 8), pady=(14, 4)
        )
        ctk.CTkLabel(helper, text="这是第几教学周", text_color=COLORS["muted"]).grid(
            row=0, column=2, sticky="w", padx=(16, 8), pady=(14, 4)
        )
        reference_entry = self._entry(helper, self.reference_date)
        reference_entry.grid(row=1, column=0, columnspan=2, sticky="ew", padx=(16, 8), pady=(0, 10))
        week_entry = self._entry(helper, self.reference_week)
        week_entry.grid(row=1, column=2, sticky="ew", padx=(8, 8), pady=(0, 10))
        reference_entry.bind("<KeyRelease>", lambda _event: self._update_reference_hint())
        week_entry.bind("<KeyRelease>", lambda _event: self._update_reference_hint())
        ctk.CTkButton(
            helper,
            text="自动计算并填入",
            width=140,
            height=42,
            fg_color=ACCENT,
            hover_color=ACCENT_HOVER,
            command=self.calculate_and_fill_first_week,
        ).grid(row=1, column=3, padx=(8, 16), pady=(0, 10))
        ctk.CTkLabel(
            helper,
            textvariable=self.first_week_hint,
            anchor="w",
            text_color=ACCENT,
            font=ctk.CTkFont("Segoe UI", 12, "bold"),
        ).grid(row=2, column=0, columnspan=4, sticky="ew", padx=16, pady=(0, 14))
        sections_card = self._card(page)
        sections_card.pack(fill="x", pady=(0, 16))
        toolbar = ctk.CTkFrame(sections_card, fg_color="transparent")
        toolbar.pack(fill="x", padx=20, pady=(18, 10))
        self.section_count_label = ctk.CTkLabel(
            toolbar,
            text="12 个节次",
            text_color=COLORS["text"],
            font=ctk.CTkFont("Segoe UI", 16, "bold"),
        )
        self.section_count_label.pack(anchor="w")
        toolbar_buttons = ctk.CTkFrame(toolbar, fg_color="transparent")
        toolbar_buttons.pack(fill="x", pady=(8, 0))
        for text, command in (
            ("删除所选", self.delete_selected_sections),
            ("批量平移", self.shift_selected_sections),
            ("全选/取消", self.toggle_all_sections),
            ("＋ 添加节次", self.add_section),
        ):
            primary = text.startswith("＋")
            ctk.CTkButton(
                toolbar_buttons,
                text=text,
                width=100,
                height=34,
                corner_radius=11,
                fg_color=ACCENT if primary else COLORS["soft"],
                hover_color=ACCENT_HOVER if primary else COLORS["border"],
                text_color="#FFFFFF" if primary else COLORS["text"],
                command=command,
            ).pack(side="right", padx=4)
        self.section_rows = ctk.CTkFrame(sections_card, fg_color="transparent")
        self.section_rows.pack(fill="x", padx=20, pady=(0, 16))
        ctk.CTkButton(
            sections_card,
            text="恢复默认 12 节示例时间",
            width=190,
            height=34,
            fg_color="transparent",
            border_width=1,
            border_color=COLORS["border"],
            text_color=COLORS["muted"],
            hover_color=COLORS["soft"],
            command=self._load_default_section_times,
        ).pack(anchor="w", padx=22, pady=(0, 20))
        ctk.CTkButton(
            page,
            text="下一步：生成文件  →",
            width=210,
            height=44,
            fg_color=ACCENT,
            hover_color=ACCENT_HOVER,
            command=lambda: self.navigate("export"),
        ).pack(anchor="e", pady=(0, 18))

    def _build_export_page(self) -> None:
        page = self._new_page("export", scrollable=True)
        self._section_title(
            page,
            "生成并查看输出",
            "先保存当前可编辑 Excel，再从同一个 Excel 生成 Apple Calendar 文件。",
        )
        output_card = self._card(page)
        output_card.pack(fill="x", pady=(0, 14))
        output_card.grid_columnconfigure(0, weight=1)
        ctk.CTkLabel(output_card, text="输出文件夹", anchor="w").grid(
            row=0, column=0, sticky="ew", padx=22, pady=(20, 5)
        )
        path_row = ctk.CTkFrame(output_card, fg_color="transparent")
        path_row.grid(row=1, column=0, sticky="ew", padx=22)
        path_row.grid_columnconfigure(0, weight=1)
        self._entry(path_row, self.output_dir).grid(row=0, column=0, sticky="ew", padx=(0, 10))
        ctk.CTkButton(
            path_row,
            text="选择文件夹",
            width=110,
            height=42,
            fg_color=COLORS["soft"],
            hover_color=COLORS["border"],
            text_color=COLORS["text"],
            command=self.select_output_dir,
        ).grid(row=0, column=1)
        ctk.CTkLabel(
            output_card,
            text="将创建 timetable.xlsx 和 timetable.ics",
            anchor="w",
            text_color=COLORS["muted"],
            font=ctk.CTkFont("Segoe UI", 11),
        ).grid(row=2, column=0, sticky="ew", padx=22, pady=(7, 16))
        export_actions = ctk.CTkFrame(output_card, fg_color="transparent")
        export_actions.grid(row=3, column=0, sticky="ew", padx=22, pady=(0, 22))
        self.save_excel_button = ctk.CTkButton(
            export_actions,
            text="仅保存 Excel",
            width=150,
            height=48,
            fg_color=COLORS["soft"],
            hover_color=COLORS["border"],
            text_color=COLORS["text"],
            command=lambda: self.generate_files(include_ics=False),
        )
        self.save_excel_button.pack(side="left")
        self.generate_button = AppButton(
            export_actions,
            text="生成 Excel + ICS",
            width=220,
            height=48,
            fg_color=ACCENT,
            hover_color=ACCENT_HOVER,
            command=lambda: self.generate_files(include_ics=True),
        )
        self.generate_button.pack(side="left", padx=12)
        result_card = self._card(page)
        result_card.pack(fill="x", pady=(0, 18))
        ctk.CTkLabel(
            result_card,
            text="完成情况",
            anchor="w",
            text_color=COLORS["text"],
            font=ctk.CTkFont("Segoe UI", 17, "bold"),
        ).pack(fill="x", padx=22, pady=(20, 8))
        ctk.CTkLabel(
            result_card,
            textvariable=self.completion_text,
            anchor="w",
            justify="left",
            wraplength=760,
            text_color=COLORS["muted"],
            font=ctk.CTkFont("Segoe UI", 13),
        ).pack(fill="x", padx=22, pady=(0, 16))
        open_actions = ctk.CTkFrame(result_card, fg_color="transparent")
        open_actions.pack(fill="x", padx=22, pady=(0, 22))
        self.open_folder_button = ctk.CTkButton(
            open_actions,
            text="打开 output 文件夹",
            fg_color=GREEN,
            hover_color="#2EAD4F",
            command=lambda: self._open_path(Path(self.output_dir.get())),
        )
        self.open_excel_button = ctk.CTkButton(
            open_actions,
            text="打开 Excel",
            fg_color=COLORS["soft"],
            hover_color=COLORS["border"],
            text_color=COLORS["text"],
            command=lambda: self._open_path(self.current_excel_path),
        )
        self.open_ics_button = ctk.CTkButton(
            open_actions,
            text="打开 ICS",
            fg_color=COLORS["soft"],
            hover_color=COLORS["border"],
            text_color=COLORS["text"],
            command=lambda: self._open_path(self.current_ics_path),
        )
        for button in (self.open_folder_button, self.open_excel_button, self.open_ics_button):
            button.pack(side="left", padx=(0, 10))
        self._refresh_output_buttons()

    def _build_settings_page(self) -> None:
        page = self._new_page("settings", scrollable=True)
        self._section_title(
            page,
            "外观与磨砂",
            "设置会实时生效并自动保存。Windows 11 优先使用系统 Acrylic 效果。",
        )
        card = self._card(page)
        card.pack(fill="x", pady=(0, 14))
        card.grid_columnconfigure(1, weight=1)
        rows = (
            ("明暗风格", "浅色、深色或跟随 Windows", 0),
            ("透明磨砂", "关闭后恢复完全不透明的普通窗口", 1),
            ("窗口不透明度", "数值越低，背景透过越明显", 2),
            ("磨砂强度", "低、中、高档分别使用 Mica、Mica Alt 与 Acrylic", 3),
        )
        for title, subtitle, row in rows:
            ctk.CTkLabel(
                card,
                text=title,
                anchor="w",
                text_color=COLORS["text"],
                font=ctk.CTkFont("Segoe UI", 14, "bold"),
            ).grid(row=row * 2, column=0, sticky="w", padx=24, pady=(20 if row == 0 else 12, 0))
            ctk.CTkLabel(
                card,
                text=subtitle,
                anchor="w",
                text_color=COLORS["muted"],
                font=ctk.CTkFont("Segoe UI", 11),
            ).grid(row=row * 2 + 1, column=0, sticky="w", padx=24, pady=(1, 10))
        ctk.CTkSegmentedButton(
            card,
            values=["跟随系统", "浅色", "深色"],
            variable=self.appearance_var,
            selected_color=ACCENT,
            selected_hover_color=ACCENT_HOVER,
            unselected_color=COLORS["soft"],
            unselected_hover_color=COLORS["border"],
            text_color=COLORS["text"],
            command=lambda _value: self.apply_preferences(),
        ).grid(row=0, rowspan=2, column=1, sticky="e", padx=24)
        ctk.CTkSwitch(
            card,
            text="",
            variable=self.glass_var,
            width=54,
            progress_color=ACCENT,
            command=self.apply_preferences,
        ).grid(row=2, rowspan=2, column=1, sticky="e", padx=24)
        opacity_control = ctk.CTkFrame(card, fg_color="transparent")
        opacity_control.grid(row=4, rowspan=2, column=1, sticky="e", padx=24)
        ctk.CTkSlider(
            opacity_control,
            from_=80,
            to=100,
            number_of_steps=20,
            width=250,
            variable=self.opacity_var,
            progress_color=ACCENT,
            button_color=ACCENT,
            button_hover_color=ACCENT_HOVER,
            command=lambda _value: self._on_glass_controls_changed(),
        ).pack(side="left")
        ctk.CTkLabel(
            opacity_control,
            textvariable=self.opacity_label,
            width=48,
            text_color=ACCENT,
            font=ctk.CTkFont("Segoe UI", 12, "bold"),
        ).pack(side="left", padx=(10, 0))
        blur_control = ctk.CTkFrame(card, fg_color="transparent")
        blur_control.grid(row=6, rowspan=2, column=1, sticky="e", padx=24)
        ctk.CTkSlider(
            blur_control,
            from_=0,
            to=100,
            number_of_steps=20,
            width=250,
            variable=self.blur_var,
            progress_color=ACCENT,
            button_color=ACCENT,
            button_hover_color=ACCENT_HOVER,
            command=lambda _value: self._on_glass_controls_changed(),
        ).pack(side="left")
        ctk.CTkLabel(
            blur_control,
            textvariable=self.blur_label,
            width=48,
            text_color=ACCENT,
            font=ctk.CTkFont("Segoe UI", 12, "bold"),
        ).pack(side="left", padx=(10, 0))
        ctk.CTkLabel(
            card,
            text="提示：磨砂关闭时，两个滑杆的数值会保留，但暂不作用于窗口。",
            anchor="w",
            text_color=COLORS["muted"],
            font=ctk.CTkFont("Segoe UI", 11),
        ).grid(row=8, column=0, columnspan=2, sticky="ew", padx=24, pady=(8, 20))

    def _configure_tree_style(self) -> None:
        style = ttk.Style(self.root)
        try:
            style.theme_use("clam")
        except tk.TclError:
            pass
        dark = ctk.get_appearance_mode() == "Dark"
        background = "#162338" if dark else "#FFFFFF"
        alternate = "#1A2942" if dark else "#F5F9FF"
        foreground = "#F5F8FF" if dark else "#15243B"
        heading = "#1F3554" if dark else "#E8F3FF"
        style.configure(
            "Timetable.Treeview",
            background=background,
            fieldbackground=background,
            foreground=foreground,
            borderwidth=0,
            relief="flat",
            rowheight=32,
            font=("Segoe UI", 10),
        )
        style.map(
            "Timetable.Treeview",
            background=[("selected", ACCENT)],
            foreground=[("selected", "#FFFFFF")],
        )
        style.configure(
            "Timetable.Treeview.Heading",
            background=heading,
            foreground=foreground,
            borderwidth=0,
            relief="flat",
            padding=(8, 9),
            font=("Segoe UI", 10, "bold"),
        )
        style.map("Timetable.Treeview.Heading", background=[("active", heading)])
        if hasattr(self, "course_tree"):
            self.course_tree.tag_configure("even", background=background, foreground=foreground)
            self.course_tree.tag_configure("odd", background=alternate, foreground=foreground)

    def enter_workspace(self, page: str) -> None:
        if self.in_workspace:
            self.navigate(page)
            return
        self.navigate(page)
        self.home.place_forget()
        self.workspace.place(relx=0, rely=0, relwidth=1, relheight=1)
        self.in_workspace = True

    def exit_workspace(self) -> None:
        if not self.in_workspace:
            return
        self.workspace.place_forget()
        self.home.place(relx=0, rely=0, relwidth=1, relheight=1)
        self.in_workspace = False

    def navigate(self, page: str) -> None:
        if page not in self.pages:
            return
        self.current_page = page
        # CTkScrollableFrame owns an internal canvas that does not reliably obey
        # sibling tkraise() ordering. Explicitly hide inactive pages instead.
        for name, frame in self.page_containers.items():
            if name == page:
                frame.grid()
                frame.tkraise()
            else:
                frame.grid_remove()
        self.page_title.configure(text=self.PAGE_TITLES[page])
        for name, button in self.nav_buttons.items():
            if name == "home":
                continue
            button.configure(
                fg_color=ACCENT if name == page else "transparent",
                text_color="#FFFFFF" if name == page else COLORS["text"],
            )
        if page == "export":
            self._refresh_output_buttons()

    def select_image(self) -> None:
        selected = filedialog.askopenfilename(
            parent=self.root,
            title="选择课表图片",
            filetypes=[
                ("图片文件", "*.png *.jpg *.jpeg *.webp *.gif"),
                ("所有文件", "*.*"),
            ],
        )
        if selected:
            self._set_image_path(Path(selected))

    def _on_drop(self, event: Any) -> str | None:
        paths = self.root.tk.splitlist(event.data)
        if paths:
            self._set_image_path(Path(paths[0]))
        return getattr(event, "action", None)

    def _set_image_path(self, path: Path) -> None:
        allowed = {".png", ".jpg", ".jpeg", ".webp", ".gif"}
        if path.suffix.lower() not in allowed or not path.is_file():
            messagebox.showerror("不支持的文件", "请选择 PNG、JPEG、WEBP 或 GIF 图片。")
            return
        resolved = path.resolve()
        self.image_path.set(str(resolved))
        self.drop_icon.configure(text="✓", text_color=GREEN)
        self.drop_title.configure(text=resolved.name)
        self.drop_detail.configure(text="图片已就绪\n点击此区域可以重新选择")
        self.status.set("图片已就绪。输入 API Key 后即可开始识别。")

    def recognize_image(self) -> None:
        if self.busy:
            return
        image = Path(self.image_path.get()) if self.image_path.get() else None
        if not image or not image.is_file():
            messagebox.showwarning("尚未选择图片", "请拖入或选择一张课表图片。")
            return
        api_key = self.api_key.get().strip()
        if not api_key:
            messagebox.showwarning("缺少 API Key", "请输入 OpenAI API Key。它不会被保存。")
            return
        model = self.model.get().strip() or DEFAULT_MODEL
        self._set_busy(True, "正在理解整张课表，窗口仍可移动……")

        def work() -> None:
            try:
                raw = request_timetable_json(
                    image, DEFAULT_RAW_RESPONSE, api_key=api_key, model=model
                )
                timetable = validate_and_normalize(raw)
                DEFAULT_NORMALIZED_JSON.parent.mkdir(parents=True, exist_ok=True)
                DEFAULT_NORMALIZED_JSON.write_text(
                    json.dumps(timetable, ensure_ascii=False, indent=2), encoding="utf-8"
                )
                self.root.after(0, lambda: self._recognition_finished(timetable["courses"]))
            except Exception as exc:
                self.root.after(0, lambda error=exc: self._task_failed("识别失败", error))

        threading.Thread(target=work, daemon=True).start()

    def _recognition_finished(self, courses: list[dict[str, Any]]) -> None:
        self.courses = courses
        self._refresh_course_tree()
        self._set_busy(False, f"识别完成，共 {len(courses)} 门课程。请逐行核对。")
        self.navigate("courses")
        if not courses:
            messagebox.showwarning("未识别到课程", "模型没有返回课程，请检查图片是否完整清晰。")

    def add_course(self) -> None:
        dialog = CourseDialog(self.root)
        if dialog.result:
            self.courses.append(dialog.result)
            self._refresh_course_tree(select_indices=[len(self.courses) - 1])

    def _selected_indices(self, require: bool = True) -> list[int]:
        indices = sorted(int(item) for item in self.course_tree.selection())
        if require and not indices:
            messagebox.showinfo("请选择课程", "请先选择一门或多门课程。")
        return indices

    def edit_selected_course(self) -> None:
        indices = self._selected_indices()
        if not indices:
            return
        if len(indices) > 1:
            self.batch_edit_courses()
            return
        index = indices[0]
        dialog = CourseDialog(self.root, self.courses[index])
        if dialog.result:
            self.courses[index] = dialog.result
            self._refresh_course_tree(select_indices=[index])

    def batch_edit_courses(self) -> None:
        indices = self._selected_indices()
        if not indices:
            return
        dialog = BatchCourseDialog(self.root, len(indices))
        if not dialog.result:
            return
        try:
            self.courses = apply_course_patch(self.courses, indices, dialog.result)
            validate_courses_for_calendar(self.courses)
            self._refresh_course_tree(select_indices=indices)
            self.status.set(f"已批量修改 {len(indices)} 门课程。")
        except (TimetableValidationError, TypeError, ValueError) as exc:
            messagebox.showerror("批量修改失败", str(exc))

    def delete_selected_courses(self) -> None:
        indices = self._selected_indices()
        if not indices:
            return
        if not messagebox.askyesno(
            "删除课程", f"确定删除选中的 {len(indices)} 门课程吗？", parent=self.root
        ):
            return
        for index in reversed(indices):
            del self.courses[index]
        self._refresh_course_tree()
        self.status.set(f"已删除 {len(indices)} 门课程。")

    def _refresh_course_tree(self, select_indices: list[int] | None = None) -> None:
        for item in self.course_tree.get_children():
            self.course_tree.delete(item)
        for index, course in enumerate(self.courses):
            values = (
                course["course_name"],
                course["weekday"],
                f"{course['start_section']}-{course['end_section']}",
                f"{course['start_week']}-{course['end_week']}",
                course["week_pattern"],
                ",".join(map(str, course["custom_weeks"])),
                course["location"],
                course["teacher"],
                course["notes"],
            )
            self.course_tree.insert(
                "", "end", iid=str(index), values=values, tags=("even" if index % 2 == 0 else "odd",)
            )
        self.course_count.set(f"{len(self.courses)} 门课程")
        if select_indices:
            valid = [str(index) for index in select_indices if 0 <= index < len(self.courses)]
            if valid:
                self.course_tree.selection_set(valid)
                self.course_tree.focus(valid[0])
                self.course_tree.see(valid[0])

    def _update_reference_hint(self) -> None:
        try:
            reference = parse_date(self.reference_date.get(), "参考日期")
            week = int(self.reference_week.get().strip())
            first = calculate_first_week_monday(reference, week)
            self.first_week_hint.set(f"计算结果：第 1 教学周周一是 {first.isoformat()}")
        except (TypeError, ValueError):
            self.first_week_hint.set("填写一个周一日期和对应教学周，例如 2026-09-14、第 3 周")

    def calculate_and_fill_first_week(self) -> None:
        try:
            reference = parse_date(self.reference_date.get(), "参考日期")
            week = int(self.reference_week.get().strip())
            first = calculate_first_week_monday(reference, week)
            self.first_week_monday.set(first.isoformat())
            self._update_reference_hint()
            self.status.set(f"已把第一教学周周一设置为 {first.isoformat()}。")
        except (TypeError, ValueError) as exc:
            messagebox.showerror("无法计算", str(exc))

    def _load_default_section_times(self) -> None:
        settings = dict(default_settings())
        values = [
            (
                section,
                format_time_value(settings[f"section_{section}_start"]),
                format_time_value(settings[f"section_{section}_end"]),
            )
            for section in range(1, 13)
        ]
        self._set_section_rows(values)

    def _section_snapshot(self) -> list[tuple[int, str, str]]:
        return [
            (section, start.get(), end.get())
            for section, (_selected, start, end) in sorted(self.section_variables.items())
        ]

    def _set_section_rows(self, values: list[tuple[int, str, str]]) -> None:
        for widget in self.section_rows.winfo_children():
            widget.destroy()
        self.section_variables.clear()
        for column, weight in enumerate((0, 0, 1, 1)):
            self.section_rows.grid_columnconfigure(column, weight=weight)
        for column, header in enumerate(("选择", "节次", "开始时间", "结束时间")):
            ctk.CTkLabel(
                self.section_rows,
                text=header,
                text_color=COLORS["muted"],
                font=ctk.CTkFont("Segoe UI", 12, "bold"),
            ).grid(row=0, column=column, sticky="w", padx=10, pady=(4, 8))
        for display_index, (_old_section, start_text, end_text) in enumerate(values, start=1):
            selected = tk.BooleanVar(value=False)
            start = tk.StringVar(value=start_text)
            end = tk.StringVar(value=end_text)
            self.section_variables[display_index] = (selected, start, end)
            ctk.CTkCheckBox(
                self.section_rows,
                text="",
                variable=selected,
                width=24,
                checkbox_width=21,
                checkbox_height=21,
                fg_color=ACCENT,
            ).grid(row=display_index, column=0, padx=10, pady=6)
            ctk.CTkLabel(
                self.section_rows,
                text=f"第 {display_index} 节",
                text_color=COLORS["text"],
                font=ctk.CTkFont("Segoe UI", 13, "bold"),
            ).grid(row=display_index, column=1, sticky="w", padx=10, pady=6)
            self._entry(self.section_rows, start).grid(
                row=display_index, column=2, sticky="ew", padx=10, pady=6
            )
            self._entry(self.section_rows, end).grid(
                row=display_index, column=3, sticky="ew", padx=10, pady=6
            )
        self.section_count_label.configure(text=f"{len(values)} 个节次")

    def add_section(self) -> None:
        values = self._section_snapshot()
        start_text = ""
        end_text = ""
        if values:
            try:
                previous_end = parse_time(values[-1][2], "上一节结束时间")
                base = datetime.combine(date(2000, 1, 1), previous_end)
                candidate_start = base + timedelta(minutes=10)
                candidate_end = candidate_start + timedelta(minutes=45)
                if candidate_end.date() == base.date():
                    start_text = candidate_start.strftime("%H:%M")
                    end_text = candidate_end.strftime("%H:%M")
            except ValueError:
                pass
        values.append((len(values) + 1, start_text, end_text))
        self._set_section_rows(values)
        self.status.set(f"已添加第 {len(values)} 节，请确认时间。")

    def toggle_all_sections(self) -> None:
        should_select = not all(
            selected.get() for selected, _start, _end in self.section_variables.values()
        )
        for selected, _start, _end in self.section_variables.values():
            selected.set(should_select)

    def delete_selected_sections(self) -> None:
        selected_numbers = [
            section
            for section, (selected, _start, _end) in self.section_variables.items()
            if selected.get()
        ]
        if not selected_numbers:
            messagebox.showinfo("请选择节次", "请先勾选要删除的一节或多节。")
            return
        if len(selected_numbers) == len(self.section_variables):
            messagebox.showwarning("不能全部删除", "至少需要保留一个节次。")
            return
        if not messagebox.askyesno(
            "删除节次",
            f"删除 {len(selected_numbers)} 个节次后，其余节次会重新连续编号。是否继续？",
            parent=self.root,
        ):
            return
        values = [item for item in self._section_snapshot() if item[0] not in selected_numbers]
        self._set_section_rows(values)
        self.status.set(f"已删除 {len(selected_numbers)} 个节次。请检查课程引用的节次。")

    def shift_selected_sections(self) -> None:
        selected_numbers = [
            section
            for section, (selected, _start, _end) in self.section_variables.items()
            if selected.get()
        ]
        if not selected_numbers:
            messagebox.showinfo("请选择节次", "请先勾选需要整体平移时间的节次。")
            return
        minutes = simpledialog.askinteger(
            "批量平移时间",
            "输入分钟数。正数向后，负数向前，例如 10 或 -5：",
            parent=self.root,
            minvalue=-720,
            maxvalue=720,
        )
        if minutes is None:
            return
        try:
            for section in selected_numbers:
                _selected, start_var, end_var = self.section_variables[section]
                day = date(2000, 1, 1)
                start_dt = datetime.combine(day, parse_time(start_var.get(), f"第 {section} 节开始时间"))
                end_dt = datetime.combine(day, parse_time(end_var.get(), f"第 {section} 节结束时间"))
                shifted_start = start_dt + timedelta(minutes=minutes)
                shifted_end = end_dt + timedelta(minutes=minutes)
                if shifted_start.date() != day or shifted_end.date() != day:
                    raise ValueError("平移后时间跨越了当天 00:00，请缩小分钟数")
                start_var.set(shifted_start.strftime("%H:%M"))
                end_var.set(shifted_end.strftime("%H:%M"))
            self.status.set(f"已将 {len(selected_numbers)} 个节次整体平移 {minutes} 分钟。")
        except ValueError as exc:
            messagebox.showerror("无法平移", str(exc))

    def _collect_settings(self) -> dict[str, Any]:
        first_monday = parse_date(self.first_week_monday.get(), "第一教学周周一")
        if first_monday.weekday() != 0:
            raise ValueError("第一教学周周一必须确实是周一")
        semester_end = parse_date(self.semester_end.get(), "学期结束日期")
        if semester_end < first_monday:
            raise ValueError("学期结束日期不能早于第一教学周")
        try:
            alarm = int(self.alarm_minutes.get().strip())
        except ValueError as exc:
            raise ValueError("课前提醒分钟数必须是整数") from exc
        if alarm < 0:
            raise ValueError("课前提醒分钟数不能小于 0")
        timezone_name = self.timezone.get().strip() or "Asia/Shanghai"
        try:
            ZoneInfo(timezone_name)
        except ZoneInfoNotFoundError as exc:
            raise ValueError(f"无法识别时区：{timezone_name}") from exc
        settings: dict[str, Any] = {
            "semester_name": self.semester_name.get().strip() or "课程表",
            "first_week_monday": first_monday,
            "semester_end": semester_end,
            "default_alarm_minutes": alarm,
            "timezone": timezone_name,
        }
        previous_end: time | None = None
        for section, (_selected, start_var, end_var) in sorted(self.section_variables.items()):
            start = parse_time(start_var.get(), f"第 {section} 节开始时间")
            end = parse_time(end_var.get(), f"第 {section} 节结束时间")
            if end <= start:
                raise ValueError(f"第 {section} 节结束时间必须晚于开始时间")
            if previous_end and start < previous_end:
                raise ValueError(f"第 {section} 节开始时间不能早于上一节结束时间")
            settings[f"section_{section}_start"] = start
            settings[f"section_{section}_end"] = end
            previous_end = end
        return settings

    def load_existing_excel(self) -> None:
        selected = filedialog.askopenfilename(
            parent=self.root,
            title="选择已有课表 Excel",
            initialdir=str(DEFAULT_OUTPUT_DIR),
            filetypes=[("Excel 工作簿", "*.xlsx")],
        )
        if not selected:
            return
        try:
            path = Path(selected)
            courses, settings = read_excel_for_gui(path)
            self.courses = courses
            self.semester_name.set(str(settings.get("semester_name", self.semester_name.get())))
            first_monday = settings.get("first_week_monday")
            if first_monday:
                first_text = format_date_value(first_monday)
                self.first_week_monday.set(first_text)
                self.reference_date.set(first_text)
                self.reference_week.set("1")
            if settings.get("semester_end"):
                self.semester_end.set(format_date_value(settings["semester_end"]))
            if settings.get("default_alarm_minutes") is not None:
                self.alarm_minutes.set(str(settings["default_alarm_minutes"]))
            if settings.get("timezone"):
                self.timezone.set(str(settings["timezone"]))
            section_numbers = sorted(
                {
                    int(match.group(1))
                    for key in settings
                    if (match := re.fullmatch(r"section_(\d+)_(?:start|end)", key))
                }
            )
            values = [
                (
                    section,
                    format_time_value(settings[f"section_{section}_start"]),
                    format_time_value(settings[f"section_{section}_end"]),
                )
                for section in section_numbers
                if f"section_{section}_start" in settings and f"section_{section}_end" in settings
            ]
            if values:
                self._set_section_rows(values)
            self.output_dir.set(str(path.parent))
            self._update_reference_hint()
            self._refresh_course_tree()
            self.enter_workspace("courses")
            self.status.set(f"已加载 {len(courses)} 门课程和 {len(values)} 个节次设置。")
        except Exception as exc:
            messagebox.showerror("无法加载 Excel", str(exc))

    def select_output_dir(self) -> None:
        selected = filedialog.askdirectory(
            parent=self.root,
            title="选择输出文件夹",
            initialdir=self.output_dir.get() or str(DEFAULT_OUTPUT_DIR),
        )
        if selected:
            self.output_dir.set(str(Path(selected).resolve()))
            self._refresh_output_buttons()

    def generate_files(self, *, include_ics: bool) -> None:
        if self.busy:
            return
        if not self.courses:
            messagebox.showwarning("没有课程", "请先识别图片、加载 Excel 或手动添加课程。")
            return
        try:
            settings = self._collect_settings()
            timetable = validate_and_normalize({"courses": self.courses})
            validate_courses_for_calendar(timetable["courses"], len(self.section_variables))
            output_directory = Path(self.output_dir.get().strip()).expanduser().resolve()
            if output_directory.exists() and not output_directory.is_dir():
                raise ValueError("输出路径不是文件夹")
        except (TimetableValidationError, OSError, ValueError) as exc:
            messagebox.showerror("无法生成", str(exc))
            return
        self._set_busy(True, "正在保存 Excel 和日历文件……" if include_ics else "正在保存 Excel……")

        def work() -> None:
            try:
                output_directory.mkdir(parents=True, exist_ok=True)
                excel_path = output_directory / "timetable.xlsx"
                ics_path = output_directory / "timetable.ics"
                write_excel(
                    timetable,
                    excel_path,
                    settings_overrides=settings,
                    preserve_existing_settings=False,
                    replace_settings=True,
                )
                result = convert_excel_to_ics(excel_path, ics_path) if include_ics else None
                self.root.after(
                    0,
                    lambda: self._generation_finished(
                        excel_path, ics_path if include_ics else None, result
                    ),
                )
            except Exception as exc:
                self.root.after(0, lambda error=exc: self._task_failed("生成失败", error))

        threading.Thread(target=work, daemon=True).start()

    def _generation_finished(self, excel_path: Path, ics_path: Path | None, result: Any) -> None:
        self.current_excel_path = excel_path
        if ics_path is not None:
            self.current_ics_path = ics_path
        if result is None:
            summary = "Excel 已保存，可以继续修改或稍后生成 ICS。"
            status = "完成：Excel 已保存。"
        else:
            summary = f"已生成 {result.event_count} 个日历事件。"
            status = f"完成：已生成 {result.event_count} 个日历事件。"
            if result.errors:
                summary += f"\n有 {result.skipped_course_count} 行课程被跳过，请检查提示。"
        details = [summary, f"Excel：{excel_path}"]
        if ics_path is not None:
            details.append(f"ICS：{ics_path}")
        self.completion_text.set("\n".join(details))
        self._set_busy(False, status)
        self._refresh_output_buttons()
        self.navigate("export")

    def _set_busy(self, busy: bool, status: str) -> None:
        self.busy = busy
        self.status.set(status)
        state = "disabled" if busy else "normal"
        for button in (self.recognize_button, self.generate_button, self.save_excel_button):
            button.configure(state=state)
        if busy:
            self.progress.start()
        else:
            self.progress.stop()
            self.progress.set(0)

    def _task_failed(self, title: str, error: Exception) -> None:
        self._set_busy(False, f"{title}：{error}")
        messagebox.showerror(title, str(error))

    def _refresh_output_buttons(self) -> None:
        output_directory = Path(self.output_dir.get() or DEFAULT_OUTPUT_DIR)
        if hasattr(self, "open_folder_button"):
            self.open_folder_button.configure(
                state="normal" if output_directory.exists() else "disabled"
            )
            self.open_excel_button.configure(
                state="normal" if self.current_excel_path.exists() else "disabled"
            )
            self.open_ics_button.configure(
                state="normal" if self.current_ics_path.exists() else "disabled"
            )

    def _open_path(self, path: Path) -> None:
        if not path.exists():
            messagebox.showwarning("文件不存在", f"尚未生成：{path}")
            return
        try:
            os.startfile(path)  # type: ignore[attr-defined]
        except OSError as exc:
            messagebox.showerror("无法打开", str(exc))

    def _on_glass_controls_changed(self) -> None:
        self.opacity_label.set(f"{round(self.opacity_var.get())}%")
        self.blur_label.set(f"{round(self.blur_var.get())}%")
        self.apply_preferences(announce=False)

    def apply_preferences(self, _value: Any = None, *, announce: bool = True) -> None:
        appearance = self.appearance_var.get()
        appearance_changed = self.ui_settings.get("appearance") != appearance
        ctk.set_appearance_mode(APPEARANCE_TO_CTK.get(appearance, "Light"))
        self.ui_settings = {
            "appearance": appearance,
            "glass": bool(self.glass_var.get()),
            "opacity": max(80, min(100, round(self.opacity_var.get()))),
            "blur": max(0, min(100, round(self.blur_var.get()))),
        }
        _save_ui_settings(self.ui_settings)
        if appearance_changed:
            self._configure_tree_style()
            self._refresh_course_tree()
        self._apply_window_effects()
        if announce:
            self.status.set("外观与磨砂设置已保存。")

    def _apply_window_effects(self) -> None:
        _windows_backdrop(
            self.root,
            bool(self.glass_var.get()),
            max(80, min(100, round(self.opacity_var.get()))),
            max(0, min(100, round(self.blur_var.get()))),
        )

    def _close(self) -> None:
        _save_ui_settings(
            {
                "appearance": self.appearance_var.get(),
                "glass": bool(self.glass_var.get()),
                "opacity": max(80, min(100, round(self.opacity_var.get()))),
                "blur": max(0, min(100, round(self.blur_var.get()))),
            }
        )
        self.root.destroy()


def enable_windows_dpi_awareness() -> None:
    try:
        ctypes.windll.shcore.SetProcessDpiAwareness(1)
    except (AttributeError, OSError):
        pass


def main() -> None:
    enable_windows_dpi_awareness()
    ctk.set_default_color_theme("blue")
    root = AppRoot()
    TimetableApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
